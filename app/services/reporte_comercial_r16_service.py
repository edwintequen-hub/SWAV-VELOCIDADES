"""
=========================================================
SWAV / FLOTA
Reporte Comercial R1.6 desde historico_expediciones
=========================================================

FUENTES:
- historico_expediciones: salidas comerciales reales R1.6
- servicios: asignacion Codigo TS -> Terminal
- catalogos/U8_COMERCIALES.xlsx
- catalogos/U9_COMERCIALES.xlsx

REGLAS:
- Una salida se deduplica por:
  fecha_operacional + unidad + servicio + patente + inicio_servicio
- Terminal real: unidad + codigo_ts
- Perfil: CODIGO_USUARIO + unidad -> terminal
- Perfil base: 48 periodos de 30 minutos
- Vista 1 hora: suma de dos medias horas
- Fila TOTAL del Excel no se considera servicio.
=========================================================
"""

from __future__ import annotations

from collections import defaultdict
import time as time_module
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session


ROOT = Path(__file__).resolve().parents[2]

PERFILES = {
    "U8": ROOT / "catalogos" / "U8_COMERCIALES.xlsx",
    "U9": ROOT / "catalogos" / "U9_COMERCIALES.xlsx",
}

HOJAS = {
    "LABORAL": "LAB BPMH COM",
    "SABADO": "SAB BPMH COM",
    "DOMINGO": "DOM BPMH COM",
}


def _norm(valor: Any) -> str:
    return str(valor or "").strip().upper()


def _normalizar_tipo_dia(valor: Any) -> str:
    texto = _norm(valor)

    equivalencias = {
        "LAB": "LABORAL",
        "LABORAL": "LABORAL",
        "DIA NORMAL": "LABORAL",
        "D?A NORMAL": "LABORAL",
        "NORMAL": "LABORAL",
        "SABADO": "SABADO",
        "S?BADO": "SABADO",
        "DOMINGO": "DOMINGO",
    }

    return equivalencias.get(texto, texto)


def _tipo_dia_calendario(fecha: date) -> str:
    dia = fecha.weekday()

    if dia == 5:
        return "SABADO"

    if dia == 6:
        return "DOMINGO"

    return "LABORAL"


def _parse_fecha(valor: str | date | datetime) -> date:
    if isinstance(valor, datetime):
        return valor.date()

    if isinstance(valor, date):
        return valor

    return datetime.strptime(
        str(valor),
        "%Y-%m-%d",
    ).date()


def _parse_datetime(valor: Any) -> datetime:
    if isinstance(valor, datetime):
        return valor

    return datetime.fromisoformat(str(valor))


def _hora_excel(valor: Any) -> str:
    if isinstance(valor, datetime):
        return valor.strftime("%H:%M")

    if isinstance(valor, time):
        return valor.strftime("%H:%M")

    if isinstance(valor, (int, float)):
        minutos = round(float(valor) * 24 * 60)
        hora = (minutos // 60) % 24
        minuto = minutos % 60
        return f"{hora:02d}:{minuto:02d}"

    return str(valor).strip()


def _periodo_30(dt: datetime) -> str:
    minuto = "00" if dt.minute < 30 else "30"
    return f"{dt.hour:02d}:{minuto}"


def _periodo_60_desde_30(periodo: str) -> str:
    hora = periodo.split(":", 1)[0]
    return f"{hora}:00"


def _periodos_30() -> list[str]:
    salida = []

    for hora in range(24):
        salida.append(f"{hora:02d}:00")
        salida.append(f"{hora:02d}:30")

    return salida


def _periodos_60() -> list[str]:
    return [
        f"{hora:02d}:00"
        for hora in range(24)
    ]



# =========================================================
# VARIANTES TS HISTORICAS CERTIFICADAS
#
# No modifica codigo_ts ni historico_expediciones.
# Solo permite asignar terminal a nomenclaturas historicas
# verificadas contra continuidad R1.6 + catalogo servicios.
# =========================================================

TERMINALES_TS_HISTORICOS_CERTIFICADOS = {
    ("U8", "830E"): "EL RETIRO",
    ("U9", "902E"): "AGUIRRE LUCO",
    ("U9", "923"): "PIE ANDINO",
    ("U9", "930"): "JUANITA",
    ("U9", "950E"): "JUANITA",
}


def _catalogo_terminal_ts(
    db: Session,
) -> dict[tuple[str, str], str]:

    filas = db.execute(text("""
        SELECT
            unidad,
            UPPER(TRIM(codigo_ts)) AS codigo_ts,
            COUNT(DISTINCT UPPER(TRIM(terminal))) AS cantidad_terminales,
            MAX(UPPER(TRIM(terminal))) AS terminal
        FROM servicios
        WHERE activo IS NOT FALSE
        GROUP BY
            unidad,
            UPPER(TRIM(codigo_ts))
    """)).mappings().all()

    catalogo = {}

    for fila in filas:
        if int(fila["cantidad_terminales"] or 0) != 1:
            continue

        catalogo[
            (
                _norm(fila["unidad"]),
                _norm(fila["codigo_ts"]),
            )
        ] = _norm(fila["terminal"])

    # Incorporar exclusivamente variantes historicas
    # previamente certificadas.
    # No modifica codigo_ts ni historico_expediciones.
    for clave, terminal in (
        TERMINALES_TS_HISTORICOS_CERTIFICADOS.items()
    ):
        catalogo.setdefault(
            (
                _norm(clave[0]),
                _norm(clave[1]),
            ),
            _norm(terminal),
        )

    return catalogo


def _catalogo_servicio_comercial_ts(
    db: Session,
) -> dict[tuple[str, str], str]:
    """Asocia codigo TS con servicio comercial si es unico."""
    filas = db.execute(text("""
        SELECT
            unidad,
            UPPER(TRIM(codigo_ts)) AS codigo_ts,
            COUNT(DISTINCT UPPER(TRIM(servicio))) AS cantidad_servicios,
            MAX(UPPER(TRIM(servicio))) AS servicio
        FROM servicios
        WHERE activo IS NOT FALSE
          AND codigo_ts IS NOT NULL
        GROUP BY unidad, UPPER(TRIM(codigo_ts))
    """)).mappings().all()

    catalogo = {}

    for fila in filas:
        if int(fila["cantidad_servicios"] or 0) != 1:
            continue

        catalogo[
            (_norm(fila["unidad"]), _norm(fila["codigo_ts"]))
        ] = _norm(fila["servicio"])

    return catalogo

def _catalogo_terminal_servicio(
    db: Session,
) -> dict[tuple[str, str], str]:

    filas = db.execute(text("""
        SELECT
            unidad,
            UPPER(TRIM(servicio)) AS servicio,
            COUNT(DISTINCT UPPER(TRIM(terminal))) AS cantidad_terminales,
            MAX(UPPER(TRIM(terminal))) AS terminal
        FROM servicios
        WHERE activo IS NOT FALSE
        GROUP BY
            unidad,
            UPPER(TRIM(servicio))
    """)).mappings().all()

    catalogo = {}

    for fila in filas:
        if int(fila["cantidad_terminales"] or 0) != 1:
            continue

        catalogo[
            (
                _norm(fila["unidad"]),
                _norm(fila["servicio"]),
            )
        ] = _norm(fila["terminal"])

    return catalogo


def _obtener_tipo_dia_real(
    db: Session,
    fecha: date,
    unidad: str,
) -> str:

    filas = db.execute(text("""
        SELECT
            tipo_dia,
            COUNT(*) AS cantidad
        FROM historico_expediciones
        WHERE fecha_operacional = :fecha
          AND unidad = :unidad
        GROUP BY tipo_dia
        ORDER BY cantidad DESC
    """), {
        "fecha": fecha.isoformat(),
        "unidad": unidad,
    }).mappings().all()

    for fila in filas:
        tipo = _normalizar_tipo_dia(
            fila["tipo_dia"]
        )

        if tipo in HOJAS:
            return tipo

    return _tipo_dia_calendario(fecha)


def _cargar_perfil_30(
    db: Session,
    unidad: str,
    tipo_dia: str,
) -> tuple[
    dict[tuple[str, str], int],
    dict[tuple[str, str, str], int],
    list[str],
]:

    archivo = PERFILES[unidad]

    if not archivo.exists():
        raise FileNotFoundError(
            f"No existe perfil comercial: {archivo}"
        )

    hoja = HOJAS[tipo_dia]

    catalogo = _catalogo_terminal_servicio(db)

    wb = load_workbook(
        archivo,
        data_only=True,
        read_only=True,
    )

    try:
        ws = wb[hoja]

        columnas = {}

        for columna in range(2, 50):
            columnas[columna] = _hora_excel(
                ws.cell(1, columna).value
            )

        por_terminal = defaultdict(int)
        por_servicio = defaultdict(int)
        faltantes = []

        for fila in range(2, ws.max_row + 1):

            servicio = _norm(
                ws.cell(fila, 1).value
            )

            if not servicio:
                continue

            if servicio == "TOTAL":
                continue

            terminal = catalogo.get(
                (unidad, servicio)
            )

            if not terminal:
                faltantes.append(servicio)
                continue

            for columna, periodo in columnas.items():

                valor = ws.cell(
                    fila,
                    columna,
                ).value

                try:
                    cantidad = int(valor or 0)
                except (TypeError, ValueError):
                    cantidad = 0

                por_terminal[
                    (terminal, periodo)
                ] += cantidad

                por_servicio[
                    (
                        terminal,
                        servicio,
                        periodo,
                    )
                ] += cantidad

        return (
            dict(por_terminal),
            dict(por_servicio),
            sorted(set(faltantes)),
        )

    finally:
        wb.close()


def _cargar_real_30(
    db: Session,
    fecha: date,
    unidad: str,
    incluir_detalle: bool = True,
) -> tuple[
    dict[tuple[str, str], int],
    dict[tuple[str, str, str], int],
    list[dict],
    datetime | None,
    datetime | None,
    list[str],
]:

    catalogo = _catalogo_terminal_ts(db)
    catalogo_comercial = _catalogo_servicio_comercial_ts(db)

    filas = db.execute(text("""
        SELECT
            fecha_operacional,
            unidad,
            servicio,
            patente,
            inicio_servicio,
            MAX(codigo_ts) AS codigo_ts,
            MAX(ruta_normalizada) AS ruta_normalizada,
            MAX(sentido) AS sentido
        FROM historico_expediciones
        WHERE fecha_operacional = :fecha
          AND unidad = :unidad
        GROUP BY
            fecha_operacional,
            unidad,
            servicio,
            patente,
            inicio_servicio
        ORDER BY inicio_servicio
    """), {
        "fecha": fecha.isoformat(),
        "unidad": unidad,
    }).mappings().all()

    por_terminal = defaultdict(int)
    por_servicio = defaultdict(int)

    detalle = []
    sin_terminal = []

    minimo = None
    maximo = None

    for fila in filas:

        codigo_ts = _norm(
            fila["codigo_ts"]
        )

        terminal = catalogo.get(
            (unidad, codigo_ts)
        )

        if not terminal:
            sin_terminal.append(codigo_ts)
            continue

        dt = _parse_datetime(
            fila["inicio_servicio"]
        )

        periodo = _periodo_30(dt)
        servicio_original = _norm(fila["servicio"])
        servicio = catalogo_comercial.get(
            (unidad, codigo_ts),
            servicio_original,
        )
        patente = _norm(fila["patente"])

        por_terminal[
            (terminal, periodo)
        ] += 1

        por_servicio[
            (
                terminal,
                servicio,
                periodo,
            )
        ] += 1

        if incluir_detalle:
            detalle.append({
                "unidad": unidad,
                "terminal": terminal,
                "servicio": servicio,
                "servicio_r16_original": servicio_original,
                "codigo_ts": codigo_ts,
                "patente": patente,
                "inicio_servicio": dt.isoformat(
                    sep=" "
                ),
                "periodo_30": periodo,
                "periodo_1h":
                    _periodo_60_desde_30(periodo),
                "ruta":
                    fila["ruta_normalizada"],
                "sentido":
                    fila["sentido"],
            })

        if minimo is None or dt < minimo:
            minimo = dt

        if maximo is None or dt > maximo:
            maximo = dt

    return (
        dict(por_terminal),
        dict(por_servicio),
        detalle,
        minimo,
        maximo,
        sorted(set(sin_terminal)),
    )


def _agrupar_1h(
    datos: dict[tuple[str, str], int],
) -> dict[tuple[str, str], int]:

    resultado = defaultdict(int)

    for (terminal, periodo), cantidad in datos.items():

        resultado[
            (
                terminal,
                _periodo_60_desde_30(periodo),
            )
        ] += cantidad

    return dict(resultado)


def _agrupar_servicio_1h(
    datos: dict[tuple[str, str, str], int],
) -> dict[tuple[str, str, str], int]:

    resultado = defaultdict(int)

    for (
        terminal,
        servicio,
        periodo,
    ), cantidad in datos.items():

        resultado[
            (
                terminal,
                servicio,
                _periodo_60_desde_30(periodo),
            )
        ] += cantidad

    return dict(resultado)


def construir_reporte_comercial_r16(
    db: Session,
    fecha: str | date | datetime,
    unidad: str,
    granularidad: str = "30MIN",
    incluir_detalle: bool = True,
) -> dict:

    fecha_obj = _parse_fecha(fecha)
    unidad = _norm(unidad)

    if unidad not in PERFILES:
        raise ValueError(
            "Unidad invalida. Use U8 o U9."
        )

    granularidad = _norm(granularidad)

    if granularidad not in {
        "30MIN",
        "1H",
    }:
        raise ValueError(
            "Granularidad invalida. "
            "Use 30MIN o 1H."
        )

    _trace_t0 = time_module.monotonic()
    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} TIPO_DIA INICIO",
        flush=True,
    )

    tipo_dia = _obtener_tipo_dia_real(
        db,
        fecha_obj,
        unidad,
    )

    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} TIPO_DIA FIN "
        f"{time_module.monotonic() - _trace_t0:.2f}s "
        f"tipo={tipo_dia}",
        flush=True,
    )

    _trace_t0 = time_module.monotonic()
    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} PERFIL INICIO",
        flush=True,
    )

    (
        perfil_terminal_30,
        perfil_servicio_30,
        perfil_faltantes,
    ) = _cargar_perfil_30(
        db,
        unidad,
        tipo_dia,
    )

    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} PERFIL FIN "
        f"{time_module.monotonic() - _trace_t0:.2f}s "
        f"servicios={len(perfil_servicio_30)}",
        flush=True,
    )

    _trace_t0 = time_module.monotonic()
    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} REAL_R16 INICIO",
        flush=True,
    )

    (
        real_terminal_30,
        real_servicio_30,
        detalle,
        inicio_r16,
        corte_r16,
        real_sin_terminal,
    ) = _cargar_real_30(
        db,
        fecha_obj,
        unidad,
        incluir_detalle=incluir_detalle,
    )

    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} REAL_R16 FIN "
        f"{time_module.monotonic() - _trace_t0:.2f}s "
        f"servicios={len(real_servicio_30)} "
        f"detalle={len(detalle)}",
        flush=True,
    )

    if granularidad == "1H":

        perfil_terminal = _agrupar_1h(
            perfil_terminal_30
        )

        real_terminal = _agrupar_1h(
            real_terminal_30
        )

        perfil_servicio = _agrupar_servicio_1h(
            perfil_servicio_30
        )

        real_servicio = _agrupar_servicio_1h(
            real_servicio_30
        )

        periodos = _periodos_60()

    else:

        perfil_terminal = perfil_terminal_30
        real_terminal = real_terminal_30

        perfil_servicio = perfil_servicio_30
        real_servicio = real_servicio_30

        periodos = _periodos_30()

    terminales = sorted({
        terminal
        for terminal, _ in (
            set(perfil_terminal.keys())
            |
            set(real_terminal.keys())
        )
    })

    matriz = []

    for terminal in terminales:

        periodos_terminal = []

        total_real = 0
        total_perfil = 0

        for periodo in periodos:

            real = int(
                real_terminal.get(
                    (terminal, periodo),
                    0,
                )
            )

            perfil = int(
                perfil_terminal.get(
                    (terminal, periodo),
                    0,
                )
            )

            total_real += real
            total_perfil += perfil

            periodos_terminal.append({
                "periodo": periodo,
                "real": real,
                "perfil": perfil,
                "diferencia":
                    real - perfil,
            })

        matriz.append({
            "terminal": terminal,
            "total_real": total_real,
            "total_perfil": total_perfil,
            "diferencia":
                total_real - total_perfil,
            "periodos":
                periodos_terminal,
        })

    servicios_claves = sorted(
        set(perfil_servicio.keys())
        |
        set(real_servicio.keys())
    )

    servicios = []

    for terminal, servicio, periodo in servicios_claves:

        real = int(
            real_servicio.get(
                (terminal, servicio, periodo),
                0,
            )
        )

        perfil = int(
            perfil_servicio.get(
                (terminal, servicio, periodo),
                0,
            )
        )

        servicios.append({
            "terminal": terminal,
            "servicio": servicio,
            "periodo": periodo,
            "real": real,
            "perfil": perfil,
            "estado_perfil": (
                "CON_PERFIL"
                if (terminal, servicio, periodo) in perfil_servicio
                else "SIN_PERFIL"
            ),
            "diferencia":
                real - perfil,
        })

    total_real = sum(
        item["total_real"]
        for item in matriz
    )

    total_perfil = sum(
        item["total_perfil"]
        for item in matriz
    )

    # ---------------------------------------------------------
    # ESTADO DE COBERTURA
    #
    # No interpretamos automaticamente el ultimo registro
    # como dia completo. Exponemos el corte real para que
    # frontend/reporte no confunda ausencia de datos con cero.
    # ---------------------------------------------------------

    estado_cobertura = (
        "SIN_DATOS"
        if corte_r16 is None
        else "PARCIAL_O_POR_CERTIFICAR"
    )

    return {
        "estado": "OK",
        "fuente": "HISTORICO_R1.6",
        "fecha": fecha_obj.isoformat(),
        "unidad": unidad,
        "tipo_dia": tipo_dia,
        "granularidad": granularidad,
        "perfil_archivo":
            str(PERFILES[unidad]),
        "inicio_r16":
            (
                inicio_r16.isoformat(sep=" ")
                if inicio_r16
                else None
            ),
        "corte_r16":
            (
                corte_r16.isoformat(sep=" ")
                if corte_r16
                else None
            ),
        "estado_cobertura":
            estado_cobertura,
        "total_real":
            total_real,
        "total_perfil":
            total_perfil,
        "diferencia":
            total_real - total_perfil,
        "perfil_servicios_sin_terminal":
            perfil_faltantes,
        "r16_codigos_ts_sin_terminal":
            real_sin_terminal,
        "matriz":
            matriz,
        "servicios":
            servicios,
        "detalle":
            detalle,
    }



# =========================================================
# PERSISTENCIA HISTORICA DEL REPORTE COMERCIAL R1.6
# =========================================================

def sincronizar_historico_reporte_comercial_r16(
    db: Session,
    fecha: str | date | datetime,
    unidad: str,
    commit: bool = True,
) -> dict:
    """
    Construye el reporte comercial R1.6 de 30 minutos y
    persiste su consolidado en historico_reporte_comercial_r16.

    La operacion es idempotente:
    - inserta claves nuevas;
    - actualiza claves existentes;
    - no duplica fecha/unidad/terminal/servicio/periodo_30.

    historico_expediciones permanece intacto.
    """

    fecha_obj = _parse_fecha(fecha)
    unidad = _norm(unidad)

    _trace_t0 = time_module.monotonic()
    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} CONSTRUCTOR INICIO",
        flush=True,
    )

    reporte = construir_reporte_comercial_r16(
        db=db,
        fecha=fecha_obj,
        unidad=unidad,
        granularidad="30MIN",
        incluir_detalle=False,
    )

    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} CONSTRUCTOR FIN "
        f"{time_module.monotonic() - _trace_t0:.2f}s "
        f"filas={len(reporte['servicios'])}",
        flush=True,
    )

    fecha_iso = fecha_obj.isoformat()
    ahora = datetime.now().isoformat(
        sep=" "
    )

    _trace_persistencia = time_module.monotonic()
    _trace_total_filas = len(reporte["servicios"])

    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} PERSISTENCIA INICIO "
        f"filas={_trace_total_filas}",
        flush=True,
    )

    # ---------------------------------------------------------
    # OPTIMIZACION PRODUCCION:
    # 1 SELECT para conocer las claves ya existentes.
    # 1 executemany UPSERT para persistir todas las filas.
    #
    # Se conserva exactamente la clave historica oficial:
    # fecha_operacional + unidad + terminal + servicio + periodo_30
    # ---------------------------------------------------------

    existentes = db.execute(
        text("""
            SELECT
                terminal,
                servicio,
                periodo_30
            FROM historico_reporte_comercial_r16
            WHERE fecha_operacional = :fecha
              AND unidad = :unidad
        """),
        {
            "fecha": fecha_iso,
            "unidad": unidad,
        },
    ).all()

    claves_existentes = {
        (
            str(row[0]),
            str(row[1]),
            str(row[2]),
        )
        for row in existentes
    }

    parametros_filas = []
    insertados = 0
    actualizados = 0

    for fila in reporte["servicios"]:
        parametros = {
            "fecha": fecha_iso,
            "unidad": unidad,
            "tipo_dia": reporte["tipo_dia"],
            "terminal": fila["terminal"],
            "servicio": fila["servicio"],
            "periodo": fila["periodo"],
            "real": int(fila["real"]),
            "perfil": int(fila["perfil"]),
            "diferencia": int(fila["diferencia"]),
            "estado_perfil": fila["estado_perfil"],
            "inicio_r16": reporte["inicio_r16"],
            "corte_r16": reporte["corte_r16"],
            "estado_cobertura": reporte["estado_cobertura"],
            "fuente": "HISTORICO_R1.6",
            "fecha_actualizacion": ahora,
        }

        clave = (
            str(fila["terminal"]),
            str(fila["servicio"]),
            str(fila["periodo"]),
        )

        if clave in claves_existentes:
            actualizados += 1
        else:
            insertados += 1

        parametros_filas.append(parametros)

    if parametros_filas:
        db.execute(
            text("""
                INSERT INTO historico_reporte_comercial_r16 (
                    fecha_operacional,
                    unidad,
                    tipo_dia,
                    terminal,
                    servicio,
                    periodo_30,
                    salidas_reales,
                    perfil_comercial,
                    diferencia,
                    estado_perfil,
                    inicio_r16,
                    corte_r16,
                    estado_cobertura,
                    fuente,
                    fecha_actualizacion
                )
                VALUES (
                    :fecha,
                    :unidad,
                    :tipo_dia,
                    :terminal,
                    :servicio,
                    :periodo,
                    :real,
                    :perfil,
                    :diferencia,
                    :estado_perfil,
                    :inicio_r16,
                    :corte_r16,
                    :estado_cobertura,
                    :fuente,
                    :fecha_actualizacion
                )
                ON CONFLICT (
                    fecha_operacional,
                    unidad,
                    terminal,
                    servicio,
                    periodo_30
                )
                DO UPDATE SET
                    tipo_dia = EXCLUDED.tipo_dia,
                    salidas_reales = EXCLUDED.salidas_reales,
                    perfil_comercial = EXCLUDED.perfil_comercial,
                    diferencia = EXCLUDED.diferencia,
                    estado_perfil = EXCLUDED.estado_perfil,
                    inicio_r16 = EXCLUDED.inicio_r16,
                    corte_r16 = EXCLUDED.corte_r16,
                    estado_cobertura = EXCLUDED.estado_cobertura,
                    fuente = EXCLUDED.fuente,
                    fecha_actualizacion = EXCLUDED.fecha_actualizacion
            """),
            parametros_filas,
        )

    print(
        f"[R16 COMERCIAL TRACE] {unidad} "
        f"{fecha_obj} PERSISTENCIA FIN "
        f"{time_module.monotonic() - _trace_persistencia:.2f}s "
        f"insertados={insertados} "
        f"actualizados={actualizados}",
        flush=True,
    )

    # ---------------------------------------------------------
    # NORMALIZACION HISTORICA COMERCIAL R1.6
    # Retirar exclusivamente claves tecnicas obsoletas
    # despues del UPSERT y antes del COMMIT.
    # ---------------------------------------------------------

    claves_nuevas = {
        (
            str(fila["terminal"]),
            str(fila["servicio"]),
            str(fila["periodo"]),
        )
        for fila in reporte["servicios"]
    }

    claves_obsoletas = claves_existentes - claves_nuevas

    codigos_tecnicos_permitidos = {
        "U8": {"T830 E3"},
        "U9": {"T902 E3", "T923C", "T930 C0", "T950 E3"},
    }

    permitidos = codigos_tecnicos_permitidos.get(unidad, set())

    claves_para_retirar = [
        clave
        for clave in claves_obsoletas
        if clave[1] in permitidos
    ]

    # Normalizacion historica certificada:
    # U8 / SANTA MARGARITA / 105 (TS 805C) -> 105C.
    # Exigir igualdad exacta de salidas por periodo.
    if unidad == "U8":
        reales_105c = {
            (str(fila["terminal"]), str(fila["periodo"])): int(fila["real"])
            for fila in reporte["servicios"]
            if str(fila["terminal"]) == "SANTA MARGARITA"
            and str(fila["servicio"]) == "105C"
        }

        for clave in sorted(claves_obsoletas):
            terminal, servicio, periodo = clave

            if (terminal, servicio) != ("SANTA MARGARITA", "105"):
                continue

            fila_antigua = db.execute(
                text("""
                    SELECT salidas_reales, perfil_comercial
                    FROM historico_reporte_comercial_r16
                    WHERE fecha_operacional = :fecha
                      AND unidad = :unidad
                      AND terminal = :terminal
                      AND servicio = :servicio
                      AND periodo_30 = :periodo
                """),
                {
                    "fecha": fecha_iso,
                    "unidad": unidad,
                    "terminal": terminal,
                    "servicio": servicio,
                    "periodo": periodo,
                },
            ).one_or_none()

            if fila_antigua is None:
                raise RuntimeError(
                    f"Registro historico 105 no encontrado: {clave}"
                )

            real_antiguo = int(fila_antigua[0])
            perfil_antiguo = int(fila_antigua[1])
            real_nuevo = reales_105c.get((terminal, periodo))

            if (
                real_nuevo is None
                or real_antiguo != real_nuevo
                or perfil_antiguo != 0
            ):
                raise RuntimeError(
                    f"Normalizacion 105C no certificada: "
                    f"{fecha_iso} {clave}, "
                    f"antiguo={real_antiguo}, "
                    f"nuevo={real_nuevo}, "
                    f"perfil={perfil_antiguo}"
                )

            claves_para_retirar.append(clave)

    # Nunca eliminar claves no certificadas.
    claves_no_certificadas = claves_obsoletas - set(claves_para_retirar)

    if claves_no_certificadas:
        raise RuntimeError(
            f"{unidad} {fecha_iso}: "
            f"{len(claves_no_certificadas)} claves historicas "
            "fuera del reporte sin normalizacion certificada. "
            "Sincronizacion detenida."
        )

    for terminal, servicio, periodo in claves_para_retirar:

        fila_tecnica = db.execute(
            text("""
                SELECT
                    salidas_reales,
                    perfil_comercial
                FROM historico_reporte_comercial_r16
                WHERE fecha_operacional = :fecha
                  AND unidad = :unidad
                  AND terminal = :terminal
                  AND servicio = :servicio
                  AND periodo_30 = :periodo
            """),
            {
                "fecha": fecha_iso,
                "unidad": unidad,
                "terminal": terminal,
                "servicio": servicio,
                "periodo": periodo,
            },
        ).one_or_none()

        if fila_tecnica is None:
            raise RuntimeError(
                "Registro tecnico desaparecio durante sincronizacion"
            )

        if int(fila_tecnica[1]) != 0:
            raise RuntimeError(
                f"Perfil no nulo en clave tecnica {servicio}"
            )

        db.execute(
            text("""
                DELETE FROM historico_reporte_comercial_r16
                WHERE fecha_operacional = :fecha
                  AND unidad = :unidad
                  AND terminal = :terminal
                  AND servicio = :servicio
                  AND periodo_30 = :periodo
            """),
            {
                "fecha": fecha_iso,
                "unidad": unidad,
                "terminal": terminal,
                "servicio": servicio,
                "periodo": periodo,
            },
        )

    validacion = db.execute(
        text("""
            SELECT
                COUNT(*),
                COALESCE(SUM(salidas_reales), 0),
                COALESCE(SUM(perfil_comercial), 0)
            FROM historico_reporte_comercial_r16
            WHERE fecha_operacional = :fecha
              AND unidad = :unidad
        """),
        {
            "fecha": fecha_iso,
            "unidad": unidad,
        },
    ).one()

    esperado = (
        len(claves_nuevas),
        int(reporte["total_real"]),
        int(reporte["total_perfil"]),
    )

    if tuple(validacion) != esperado:
        raise RuntimeError(
            f"{unidad} {fecha_iso}: validacion historica "
            f"incorrecta. Actual={tuple(validacion)} "
            f"Esperado={esperado}"
        )

    print(
        f"[R16 COMERCIAL] {unidad} {fecha_iso} "
        f"claves_tecnicas_retiradas={len(claves_para_retirar)} "
        f"integridad=OK",
        flush=True,
    )

    if commit:
        db.commit()

    return {
        "estado":
            "OK",
        "fecha":
            fecha_iso,
        "unidad":
            unidad,
        "tipo_dia":
            reporte["tipo_dia"],
        "insertados":
            insertados,
        "actualizados":
            actualizados,
        "total_real":
            reporte["total_real"],
        "total_perfil":
            reporte["total_perfil"],
        "diferencia":
            reporte["diferencia"],
        "inicio_r16":
            reporte["inicio_r16"],
        "corte_r16":
            reporte["corte_r16"],
        "estado_cobertura":
            reporte["estado_cobertura"],
        "perfil_servicios_sin_terminal":
            reporte[
                "perfil_servicios_sin_terminal"
            ],
        "r16_codigos_ts_sin_terminal":
            reporte[
                "r16_codigos_ts_sin_terminal"
            ],
    }
