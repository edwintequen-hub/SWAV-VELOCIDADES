from __future__ import annotations

import hashlib
import re
from collections import Counter
from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import (
    FlotaSnapshotR001,
    FlotaSnapshotR003,
)


TERMINALES = {
    "AGUIRRE LUCO",
    "CONDELL",
    "EL RETIRO",
    "JUANITA",
    "PIE ANDINO",
    "SANTA MARGARITA",
    "SANTA MARTA",
}


def _norm(valor) -> str:
    return " ".join(
        str(valor or "").strip().upper().split()
    )


def _ppu(valor) -> str:
    return re.sub(
        r"[^A-Z0-9]",
        "",
        str(valor or "").upper(),
    )


def _es_ppu(valor) -> bool:
    return bool(
        re.fullmatch(
            r"[A-Z]{4}[0-9]{2}",
            _ppu(valor),
        )
    )


def _entero(valor):
    if valor is None or valor == "":
        return None

    if isinstance(valor, bool):
        return int(valor)

    if isinstance(valor, (int, float)):
        return int(valor)

    texto = str(valor).strip()
    texto = texto.replace(".", "").replace(",", ".")

    try:
        return int(float(texto))
    except Exception:
        return None


def _hash_bytes(datos: bytes) -> str:
    return hashlib.sha256(datos).hexdigest()


def _abrir(datos: bytes):
    return load_workbook(
        BytesIO(datos),
        data_only=True,
        read_only=False,
    )


def analizar_r001(datos: bytes) -> dict:
    wb = _abrir(datos)

    try:
        ws = wb[wb.sheetnames[0]]

        encabezado = None

        for fila in range(1, min(ws.max_row, 80) + 1):
            valores = [
                _norm(ws.cell(fila, col).value)
                for col in range(1, min(ws.max_column, 20) + 1)
            ]

            if (
                "TERMINAL" in valores
                and "TOTAL FLOTA" in valores
            ):
                encabezado = fila
                break

        if encabezado is None:
            raise ValueError(
                "R001 invalido: no se encontro encabezado "
                "TERMINAL / TOTAL FLOTA."
            )

        # R001 utiliza encabezados con celdas combinadas.
        # En el formato auditado:
        #   B = Terminal
        #   E = Total Flota
        #   H = Buses Disponibles
        #   K = Buses NO Disponibles
        #
        # No se fijan posiciones de manera ciega:
        # buscamos la celda superior izquierda que contiene
        # cada encabezado real.

        mapa = {}

        for col in range(1, ws.max_column + 1):
            valor = ws.cell(encabezado, col).value

            if valor is None:
                continue

            nombre = _norm(valor)

            if nombre:
                mapa[nombre] = col

        requeridas = {
            "TERMINAL",
            "TOTAL FLOTA",
            "BUSES DISPONIBLES",
            "BUSES NO DISPONIBLES",
        }

        faltan_columnas = sorted(
            requeridas - set(mapa)
        )

        if faltan_columnas:
            raise ValueError(
                "R001 invalido: faltan columnas: "
                + ", ".join(faltan_columnas)
            )

        col_terminal = mapa["TERMINAL"]
        col_total = mapa["TOTAL FLOTA"]
        col_disp = mapa["BUSES DISPONIBLES"]
        col_no_disp = mapa["BUSES NO DISPONIBLES"]

        filas = []

        for fila in range(encabezado + 1, ws.max_row + 1):
            terminal = _norm(
                ws.cell(fila, col_terminal).value
            )

            # La tabla principal R001 termina en la fila
            # "Totales". Debajo existen graficos/resumenes que
            # vuelven a repetir nombres de terminales y NO deben
            # interpretarse como registros de flota.
            if terminal in {"TOTALES", "TOTAL"}:
                break

            if terminal not in TERMINALES:
                continue

            total = _entero(
                ws.cell(fila, col_total).value
            )

            disponibles = (
                _entero(ws.cell(fila, col_disp).value)
                if col_disp
                else None
            )

            no_disponibles = (
                _entero(ws.cell(fila, col_no_disp).value)
                if col_no_disp
                else None
            )

            if total is None:
                raise ValueError(
                    f"R001: total invalido para {terminal}."
                )

            if (
                disponibles is not None
                and no_disponibles is not None
                and total != disponibles + no_disponibles
            ):
                raise ValueError(
                    f"R001 inconsistente en {terminal}: "
                    f"{total} != {disponibles} + {no_disponibles}"
                )

            filas.append({
                "terminal": terminal,
                "total_flota": total,
                "buses_disponibles": disponibles,
                "buses_no_disponibles": no_disponibles,
            })

        encontrados = {
            x["terminal"]
            for x in filas
        }

        faltantes = sorted(
            TERMINALES - encontrados
        )

        if faltantes:
            raise ValueError(
                "R001 incompleto. Faltan terminales: "
                + ", ".join(faltantes)
            )

        if len(filas) != len(TERMINALES):
            raise ValueError(
                "R001 contiene terminales duplicados."
            )

        return {
            "filas": filas,
            "total_flota": sum(
                x["total_flota"]
                for x in filas
            ),
            "total_disponibles": sum(
                x["buses_disponibles"] or 0
                for x in filas
            ),
            "total_no_disponibles": sum(
                x["buses_no_disponibles"] or 0
                for x in filas
            ),
            "hash": _hash_bytes(datos),
        }

    finally:
        wb.close()


def analizar_r003(datos: bytes) -> dict:
    wb = _abrir(datos)

    try:
        ws = wb[wb.sheetnames[0]]

        terminal_actual = None
        registros = {}
        duplicados = Counter()

        for fila in range(1, ws.max_row + 1):

            # R003 usa bloques de terminal.
            # Se inspeccionan las primeras columnas para encontrar
            # encabezados/etiquetas de terminal.
            for col in range(1, min(ws.max_column, 4) + 1):
                candidato = _norm(
                    ws.cell(fila, col).value
                )

                if candidato in TERMINALES:
                    terminal_actual = candidato
                    break

            # El reporte actual trae DOMINIO/PPU en columna 7.
            ppu = _ppu(
                ws.cell(fila, 7).value
            )

            if not _es_ppu(ppu):
                continue

            duplicados[ppu] += 1

            if ppu in registros:
                continue

            interno = None
            tipo_bus = None

            # Conservamos datos auxiliares sin convertirlos
            # en fuente operacional.
            for col in range(1, ws.max_column + 1):
                valor = ws.cell(fila, col).value

                if valor is None:
                    continue

                texto = str(valor).strip()

                if (
                    interno is None
                    and re.fullmatch(
                        r"CL[-]?\d+",
                        texto.upper(),
                    )
                ):
                    interno = texto

            registros[ppu] = {
                "ppu": ppu,
                "terminal": terminal_actual,
                "interno": interno,
                "tipo_bus": tipo_bus,
                "fila_origen": fila,
            }

        if not registros:
            raise ValueError(
                "R003 invalido: no se encontraron PPU."
            )

        repetidas = {
            ppu: cantidad
            for ppu, cantidad in duplicados.items()
            if cantidad > 1
        }

        return {
            "filas": list(registros.values()),
            "total_ppu": len(registros),
            "duplicadas": repetidas,
            "por_terminal": dict(
                sorted(
                    Counter(
                        x["terminal"]
                        for x in registros.values()
                    ).items(),
                    key=lambda x: str(x[0]),
                )
            ),
            "hash": _hash_bytes(datos),
        }

    finally:
        wb.close()


def guardar_snapshot(
    db: Session,
    datos_r001: bytes,
    nombre_r001: str,
    datos_r003: bytes,
    nombre_r003: str,
    fecha_reporte: datetime | None = None,
) -> dict:

    r001 = analizar_r001(datos_r001)
    r003 = analizar_r003(datos_r003)

    ahora = fecha_reporte or datetime.now()

    snapshot_id = ahora.strftime(
        "%Y%m%d_%H%M%S"
    )

    existente = (
        db.query(FlotaSnapshotR001)
        .filter(
            FlotaSnapshotR001.carga_hash
            == r001["hash"]
        )
        .first()
    )

    if existente:
        raise ValueError(
            "Este R001 ya fue cargado anteriormente."
        )

    existente = (
        db.query(FlotaSnapshotR003)
        .filter(
            FlotaSnapshotR003.carga_hash
            == r003["hash"]
        )
        .first()
    )

    if existente:
        raise ValueError(
            "Este R003 ya fue cargado anteriormente."
        )

    try:
        for fila in r001["filas"]:
            db.add(
                FlotaSnapshotR001(
                    snapshot_id=snapshot_id,
                    fecha_reporte=ahora,
                    terminal=fila["terminal"],
                    total_flota=fila["total_flota"],
                    buses_disponibles=fila["buses_disponibles"],
                    buses_no_disponibles=fila["buses_no_disponibles"],
                    archivo_origen=nombre_r001,
                    carga_hash=r001["hash"],
                )
            )

        for fila in r003["filas"]:
            db.add(
                FlotaSnapshotR003(
                    snapshot_id=snapshot_id,
                    fecha_reporte=ahora,
                    ppu=fila["ppu"],
                    terminal=fila["terminal"],
                    interno=fila["interno"],
                    tipo_bus=fila["tipo_bus"],
                    archivo_origen=nombre_r003,
                    carga_hash=r003["hash"],
                )
            )

        db.commit()

    except Exception:
        db.rollback()
        raise

    return {
        "snapshot_id": snapshot_id,
        "fecha_reporte": ahora.isoformat(),
        "r001": {
            "total_flota": r001["total_flota"],
            "total_disponibles": r001["total_disponibles"],
            "total_no_disponibles": r001["total_no_disponibles"],
            "terminales": r001["filas"],
        },
        "r003": {
            "total_ppu": r003["total_ppu"],
            "por_terminal": r003["por_terminal"],
            "duplicadas": r003["duplicadas"],
        },
    }
