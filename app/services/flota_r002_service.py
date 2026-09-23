from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


R002_DEFAULT = Path(
    r"C:\Users\lenovo\Desktop\FLOTA"
    r"\R002 - DF - NO Disponibles.xlsx"
)


def _normalizar_ppu(valor: Any) -> str:
    return re.sub(
        r"[^A-Z0-9]",
        "",
        str(valor or "").upper().strip(),
    )


def _texto(valor: Any) -> str:
    return str(valor or "").strip()


def _terminal(valor: Any) -> str:
    return _texto(valor).upper()


def cargar_r002(
    ruta: str | Path | None = None,
) -> dict:
    """
    Lee R002 - DF - NO Disponibles.

    R002 es informacion complementaria del estado actual
    de disponibilidad de flota.

    IMPORTANTE:
    - NO modifica el catalogo de flota.
    - NO modifica R1.6.
    - NO recalcula Flota Operativa.
    - NO atribuye filas 'En carga' sin PPU a buses individuales.
    """

    archivo = Path(ruta) if ruta else R002_DEFAULT

    resultado = {
        "disponible": False,
        "archivo": str(archivo),
        "fecha_generacion": None,
        "ppus": {},
        "cargas_sin_ppu": {},
        "total_ppu": 0,
        "total_cargas_sin_ppu": 0,
        "error": None,
    }

    if not archivo.exists():
        resultado["error"] = "ARCHIVO_R002_NO_ENCONTRADO"
        return resultado

    try:
        wb = load_workbook(
            archivo,
            read_only=False,
            data_only=True,
        )

        ws = wb[wb.sheetnames[0]]

        # -----------------------------------------------------
        # Fecha/hora de generacion del reporte Power BI.
        # Estructura observada:
        # fila 7 -> "Generación reporte: DD-MM-YYYY HH:MM:SS"
        # -----------------------------------------------------

        for fila in range(1, min(ws.max_row, 15) + 1):
            for columna in range(1, min(ws.max_column, 15) + 1):

                valor = ws.cell(fila, columna).value

                if not valor:
                    continue

                texto = str(valor).strip()

                if "Generación reporte:" in texto:
                    fecha_txt = texto.split(
                        "Generación reporte:",
                        1,
                    )[1].strip()

                    try:
                        fecha_dt = datetime.strptime(
                            fecha_txt,
                            "%d-%m-%Y %H:%M:%S",
                        )
                        resultado["fecha_generacion"] = (
                            fecha_dt.isoformat()
                        )
                    except ValueError:
                        resultado["fecha_generacion"] = (
                            fecha_txt
                        )

        terminal_actual = ""
        ppus = {}
        cargas = Counter()

        # -----------------------------------------------------
        # Estructura R002 confirmada:
        # B = Terminal
        # C = Interno
        # D = Dominio / PPU
        # I = Estado
        # L = SoC Alta
        # N = Hora Alta
        # -----------------------------------------------------

        for fila in range(13, ws.max_row + 1):

            valor_terminal = ws.cell(fila, 2).value
            interno = ws.cell(fila, 3).value
            dominio = ws.cell(fila, 4).value
            estado = ws.cell(fila, 9).value
            soc_alta = ws.cell(fila, 12).value
            hora_alta = ws.cell(fila, 14).value

            if (
                valor_terminal is not None
                and _texto(valor_terminal)
            ):
                terminal_actual = _terminal(
                    valor_terminal
                )

            if terminal_actual == "TERMINAL":
                terminal_fila = ""
            else:
                terminal_fila = terminal_actual

            ppu = _normalizar_ppu(dominio)
            estado_txt = _texto(estado)

            # Encabezados repetidos dentro del Excel.
            if ppu in {"DOMINIO", "PPU"}:
                continue

            if ppu:
                ppus[ppu] = {
                    "ppu": ppu,
                    "terminal_r002": terminal_fila or None,
                    "interno_r002": _texto(interno) or None,
                    "estado_r002": estado_txt or None,
                    "soc_alta": soc_alta,
                    "hora_alta": (
                        _texto(hora_alta)
                        if hora_alta is not None
                        else None
                    ),
                }

            elif estado_txt.upper() == "EN CARGA":
                cargas[
                    terminal_fila or "SIN TERMINAL"
                ] += 1

        wb.close()

        resultado["disponible"] = True
        resultado["ppus"] = ppus
        resultado["cargas_sin_ppu"] = dict(cargas)
        resultado["total_ppu"] = len(ppus)
        resultado["total_cargas_sin_ppu"] = sum(
            cargas.values()
        )

        return resultado

    except Exception as exc:
        resultado["error"] = (
            f"{type(exc).__name__}: {exc}"
        )
        return resultado


def obtener_estado_ppu(
    ppu: str,
    datos_r002: dict | None = None,
) -> dict | None:
    datos = datos_r002 or cargar_r002()

    if not datos.get("disponible"):
        return None

    clave = _normalizar_ppu(ppu)

    if not clave:
        return None

    return datos.get("ppus", {}).get(clave)
