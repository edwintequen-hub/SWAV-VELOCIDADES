from __future__ import annotations

from collections import defaultdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from sqlalchemy import text
from sqlalchemy.orm import Session


MIME_XLSX = (
    "application/"
    "vnd.openxmlformats-officedocument."
    "spreadsheetml.sheet"
)


def _periodos_30():
    salida = []

    for hora in range(24):
        salida.append(f"{hora:02d}:00")
        salida.append(f"{hora:02d}:30")

    return salida


def _periodos_1h():
    return [
        f"{hora:02d}:00"
        for hora in range(24)
    ]


def _periodo_1h(periodo_30):
    return (
        str(periodo_30)
        .split(":", 1)[0]
        + ":00"
    )


def _texto_fecha(valor):
    if valor is None:
        return None

    if hasattr(valor, "isoformat"):
        return valor.isoformat()

    return str(valor)


def _texto_datetime(valor):
    if valor is None:
        return None

    return str(valor)


def generar_excel_comercial_r16(
    db: Session,
    fecha: str,
    unidad: str,
):
    """
    Genera el reporte diario comercial R1.6.

    FUENTE PRINCIPAL:
        historico_reporte_comercial_r16

    DETALLE:
        historico_expediciones

    No modifica ninguna tabla.
    No depende del archivo R1.6 original.
    No recalcula el perfil comercial.
    """

    unidad = str(
        unidad or ""
    ).strip().upper()

    if unidad not in ("U8", "U9"):
        raise ValueError(
            "Unidad invalida. Use U8 o U9."
        )

    # ============================================================
    # SNAPSHOT COMERCIAL PERSISTIDO
    # ============================================================

    filas = db.execute(
        text("""
            SELECT
                fecha_operacional,
                unidad,
                tipo_dia,
                terminal,
                servicio,
                periodo_30,
                salidas_reales,
                perfil_comercial,
                diferencia,
                inicio_r16,
                corte_r16,
                estado_cobertura,
                fuente
            FROM historico_reporte_comercial_r16
            WHERE fecha_operacional = :fecha
              AND unidad = :unidad
            ORDER BY
                terminal,
                servicio,
                periodo_30
        """),
        {
            "fecha": fecha,
            "unidad": unidad,
        },
    ).mappings().all()

    if not filas:
        raise ValueError(
            "No existe reporte comercial R1.6 "
            f"para {unidad} en {fecha}."
        )

    primera = filas[0]

    tipo_dia = str(
        primera["tipo_dia"] or ""
    )

    inicio_r16 = primera[
        "inicio_r16"
    ]

    corte_r16 = primera[
        "corte_r16"
    ]

    estado_cobertura = str(
        primera["estado_cobertura"]
        or ""
    )

    fuente = str(
        primera["fuente"]
        or "HISTORICO_R1.6"
    )

    # ============================================================
    # ESTRUCTURAS
    # ============================================================

    terminal_30 = defaultdict(
        lambda: defaultdict(
            lambda: {
                "real": 0,
                "perfil": 0,
            }
        )
    )

    servicio_30 = defaultdict(
        lambda: defaultdict(
            lambda: {
                "real": 0,
                "perfil": 0,
            }
        )
    )

    resumen_terminal = defaultdict(
        lambda: {
            "real": 0,
            "perfil": 0,
        }
    )

    total_real = 0
    total_perfil = 0

    for fila in filas:

        terminal = str(
            fila["terminal"] or ""
        )

        servicio = str(
            fila["servicio"] or ""
        )

        periodo = str(
            fila["periodo_30"] or ""
        )

        real = int(
            fila["salidas_reales"] or 0
        )

        perfil = int(
            fila["perfil_comercial"] or 0
        )

        terminal_30[
            terminal
        ][periodo]["real"] += real

        terminal_30[
            terminal
        ][periodo]["perfil"] += perfil

        servicio_30[
            (terminal, servicio)
        ][periodo]["real"] += real

        servicio_30[
            (terminal, servicio)
        ][periodo]["perfil"] += perfil

        resumen_terminal[
            terminal
        ]["real"] += real

        resumen_terminal[
            terminal
        ]["perfil"] += perfil

        total_real += real
        total_perfil += perfil

    # ============================================================
    # DETALLE R1.6 DESDE BD
    # ============================================================

    from app.services.reporte_comercial_r16_service import (
        _catalogo_terminal_ts,
    )

    catalogo_terminal = (
        _catalogo_terminal_ts(db)
    )

    detalle_raw = db.execute(
        text("""
            SELECT
                fecha_operacional,
                unidad,
                servicio,
                codigo_ts,
                patente,
                inicio_servicio,
                fin_servicio,
                ruta_normalizada,
                sentido
            FROM historico_expediciones
            WHERE fecha_operacional = :fecha
              AND unidad = :unidad
            ORDER BY
                inicio_servicio,
                servicio,
                patente
        """),
        {
            "fecha": fecha,
            "unidad": unidad,
        },
    ).mappings().all()

    unicos = {}

    for fila in detalle_raw:

        clave = (
            str(fila["fecha_operacional"]),
            str(fila["unidad"]),
            str(fila["servicio"]),
            str(fila["patente"]),
            str(fila["inicio_servicio"]),
        )

        if clave not in unicos:
            unicos[clave] = dict(fila)

    detalle = []

    for fila in unicos.values():

        codigo_ts = str(
            fila["codigo_ts"] or ""
        ).strip().upper()

        terminal = (
            catalogo_terminal.get(
                (unidad, codigo_ts)
            )
        )

        if not terminal:
            continue

        inicio = fila[
            "inicio_servicio"
        ]

        if inicio is None:
            continue

        if hasattr(inicio, "hour"):
            hora = inicio.hour
            minuto = inicio.minute
        else:
            texto_inicio = str(inicio)

            try:
                hora_texto = (
                    texto_inicio
                    .split(" ")[1][:5]
                )

                hora = int(
                    hora_texto[0:2]
                )

                minuto = int(
                    hora_texto[3:5]
                )

            except Exception:
                continue

        periodo_30 = (
            f"{hora:02d}:"
            f"{0 if minuto < 30 else 30:02d}"
        )

        periodo_1h = (
            f"{hora:02d}:00"
        )

        detalle.append({
            "fecha":
                _texto_fecha(
                    fila["fecha_operacional"]
                ),
            "unidad":
                str(
                    fila["unidad"] or ""
                ),
            "terminal":
                terminal,
            "servicio":
                str(
                    fila["servicio"] or ""
                ).strip(),
            "codigo_ts":
                codigo_ts,
            "ppu":
                str(
                    fila["patente"] or ""
                ).strip(),
            "inicio_servicio":
                _texto_datetime(
                    fila["inicio_servicio"]
                ),
            "fin_servicio":
                _texto_datetime(
                    fila["fin_servicio"]
                ),
            "sentido":
                str(
                    fila["sentido"] or ""
                ).strip(),
            "ruta":
                str(
                    fila["ruta_normalizada"]
                    or ""
                ).strip(),
            "periodo_30":
                periodo_30,
            "periodo_1h":
                periodo_1h,
        })

    detalle.sort(
        key=lambda x: (
            x["inicio_servicio"] or "",
            x["terminal"],
            x["servicio"],
            x["ppu"],
        )
    )

    # SWAV-19C4X2-SENTIDO-GRAFICO-SERVICIOS
    # Sentido informativo obtenido exclusivamente del R1.6 real.
    # No modifica ni distribuye el Perfil Comercial.
    sentidos_servicio_30 = defaultdict(set)

    for item in detalle:
        sentido_real = str(
            item.get("sentido") or ""
        ).strip().upper()

        if sentido_real:
            sentidos_servicio_30[
                (
                    item["terminal"],
                    item["servicio"],
                    item["periodo_30"],
                )
            ].add(sentido_real)


    # ============================================================
    # WORKBOOK
    # ============================================================

    wb = Workbook()

    ws_default = wb.active
    wb.remove(ws_default)

    COLOR_AZUL = "1F4E78"
    COLOR_AZUL_CLARO = "D9EAF7"
    COLOR_GRIS = "F2F2F2"
    COLOR_BLANCO = "FFFFFF"
    COLOR_AMARILLO = "FFF2CC"
    COLOR_VERDE = "E2F0D9"
    COLOR_ROJO = "FCE4D6"
    COLOR_BORDE = "B4C6E7"

    borde = Border(
        left=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
        right=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
        top=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
        bottom=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
    )

    def titulo(
        ws,
        texto_titulo,
        ultima_columna,
    ):
        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=ultima_columna,
        )

        celda = ws.cell(1, 1)
        celda.value = texto_titulo

        celda.font = Font(
            bold=True,
            size=16,
            color=COLOR_BLANCO,
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=COLOR_AZUL,
        )

        celda.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        ws.row_dimensions[1].height = 28

    def cabecera(
        ws,
        fila,
        columnas,
    ):
        for columna, valor in enumerate(
            columnas,
            start=1,
        ):
            celda = ws.cell(
                fila,
                columna,
                valor,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_BLANCO,
            )

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_AZUL,
            )

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            celda.border = borde

    def ajustar(
        ws,
        maximo=32,
    ):
        for columna in range(
            1,
            ws.max_column + 1,
        ):
            letra = get_column_letter(
                columna
            )

            ancho = 0

            for celda in ws[letra]:
                valor = celda.value

                if valor is None:
                    continue

                ancho = max(
                    ancho,
                    len(str(valor)),
                )

            ws.column_dimensions[
                letra
            ].width = min(
                max(
                    ancho + 2,
                    10,
                ),
                maximo,
            )

    # ============================================================
    # 01 RESUMEN
    # ============================================================

    ws = wb.create_sheet(
        "RESUMEN"
    )

    titulo(
        ws,
        "REPORTE DIARIO - SALIDAS COMERCIALES R1.6",
        6,
    )

    metadata = [
        ("Fecha", fecha),
        ("Unidad", unidad),
        ("Tipo d?a", tipo_dia),
        (
            "Inicio R1.6",
            _texto_datetime(
                inicio_r16
            ),
        ),
        (
            "Corte R1.6",
            _texto_datetime(
                corte_r16
            ),
        ),
        (
            "Cobertura",
            estado_cobertura,
        ),
        ("Fuente", fuente),
    ]

    fila = 3

    for etiqueta, valor in metadata:

        ws.cell(
            fila,
            1,
            etiqueta,
        ).font = Font(
            bold=True
        )

        ws.cell(
            fila,
            2,
            valor,
        )

        fila += 1

    fila += 1

    cabecera(
        ws,
        fila,
        [
            "Terminal",
            "Salidas R1.6",
            "Perfil Comercial",
            "Diferencia",
        ],
    )

    fila += 1

    for terminal in sorted(
        resumen_terminal
    ):
        real = resumen_terminal[
            terminal
        ]["real"]

        perfil = resumen_terminal[
            terminal
        ]["perfil"]

        diferencia = (
            real - perfil
        )

        ws.append([
            terminal,
            real,
            perfil,
            diferencia,
        ])

    ws.append([
        "TOTAL",
        total_real,
        total_perfil,
        total_real - total_perfil,
    ])

    for celda in ws[
        ws.max_row
    ]:
        celda.font = Font(
            bold=True
        )

        celda.fill = PatternFill(
            "solid",
            fgColor=COLOR_AZUL_CLARO,
        )

    if (
        estado_cobertura
        != "COMPLETO"
    ):
        ws["D3"] = (
            "ADVERTENCIA: cobertura "
            "parcial o pendiente de certificar. "
            "La diferencia no se interpreta "
            "como incumplimiento."
        )

        ws["D3"].fill = PatternFill(
            "solid",
            fgColor=COLOR_AMARILLO,
        )

        ws["D3"].alignment = Alignment(
            wrap_text=True,
        )

        ws.merge_cells(
            "D3:F6"
        )

    ws.freeze_panes = "A12"
    ajustar(ws, 45)

    # ============================================================
    # MATRICES
    # ============================================================

    def crear_matriz(
        nombre,
        periodos,
        granularidad,
    ):
        ws = wb.create_sheet(
            nombre
        )

        columnas = (
            ["Terminal", "Serie"]
            + periodos
            + ["TOTAL"]
        )

        titulo(
            ws,
            (
                f"SALIDAS COMERCIALES R1.6 "
                f"- {granularidad}"
            ),
            len(columnas),
        )

        cabecera(
            ws,
            3,
            columnas,
        )

        for terminal in sorted(
            terminal_30
        ):

            for serie in (
                "Salidas R1.6",
                "Perfil Comercial",
                "Diferencia",
            ):

                valores = []

                for periodo in periodos:

                    if granularidad == "30 MIN":

                        dato = terminal_30[
                            terminal
                        ][periodo]

                        real = dato[
                            "real"
                        ]

                        perfil = dato[
                            "perfil"
                        ]

                    else:

                        real = 0
                        perfil = 0

                        for periodo_30 in (
                            periodo,
                            periodo[:3] + "30",
                        ):
                            dato = terminal_30[
                                terminal
                            ][periodo_30]

                            real += dato[
                                "real"
                            ]

                            perfil += dato[
                                "perfil"
                            ]

                    if serie == "Salidas R1.6":
                        valor = real
                    elif serie == "Perfil Comercial":
                        valor = perfil
                    else:
                        valor = (
                            real - perfil
                        )

                    valores.append(
                        valor
                    )

                ws.append(
                    [
                        terminal,
                        serie,
                    ]
                    + valores
                    + [sum(valores)]
                )

                fila_actual = (
                    ws.max_row
                )

                relleno = (
                    COLOR_VERDE
                    if serie == "Salidas R1.6"
                    else
                    COLOR_AZUL_CLARO
                    if serie == "Perfil Comercial"
                    else
                    COLOR_ROJO
                )

                for celda in ws[
                    fila_actual
                ]:
                    celda.fill = (
                        PatternFill(
                            "solid",
                            fgColor=relleno,
                        )
                    )

                    celda.border = borde

        ws.freeze_panes = "C4"
        ws.auto_filter.ref = (
            f"A3:"
            f"{get_column_letter(ws.max_column)}"
            f"{ws.max_row}"
        )

        ws.column_dimensions[
            "A"
        ].width = 22

        ws.column_dimensions[
            "B"
        ].width = 20

        for columna in range(
            3,
            ws.max_column + 1,
        ):
            ws.column_dimensions[
                get_column_letter(
                    columna
                )
            ].width = 10

    crear_matriz(
        "30 MIN",
        _periodos_30(),
        "30 MIN",
    )

    crear_matriz(
        "1 HORA",
        _periodos_1h(),
        "1 HORA",
    )

    # ============================================================
    # SERVICIOS
    # ============================================================

    ws = wb.create_sheet(
        "SERVICIOS"
    )

    columnas_servicio = [
        "Terminal",
        "Servicio",
        "Sentido R1.6",
        "Periodo 30 min",
        "Salidas R1.6",
        "Perfil Comercial",
        "Diferencia",
    ]

    titulo(
        ws,
        "DESGLOSE POR SERVICIO - 30 MIN",
        len(columnas_servicio),
    )

    cabecera(
        ws,
        3,
        columnas_servicio,
    )

    for (
        terminal,
        servicio,
    ) in sorted(
        servicio_30
    ):

        for periodo in _periodos_30():

            dato = servicio_30[
                (terminal, servicio)
            ][periodo]

            real = dato[
                "real"
            ]

            perfil = dato[
                "perfil"
            ]

            sentidos = sorted(
                sentidos_servicio_30.get(
                    (
                        terminal,
                        servicio,
                        periodo,
                    ),
                    set(),
                )
            )

            sentido_texto = (
                " / ".join(sentidos)
                if sentidos
                else "-"
            )

            ws.append([
                terminal,
                servicio,
                sentido_texto,
                periodo,
                real,
                perfil,
                real - perfil,
            ])

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = ws.dimensions
    ajustar(ws, 24)

    # ------------------------------------------------------------
    # GRAFICO COMPARATIVO
    # Perfil Comercial / Salidas R1.6 / Diferencia
    # ------------------------------------------------------------

    if ws.max_row >= 4:

        grafico = BarChart()

        grafico.type = "col"
        grafico.style = 10

        grafico.title = (
            "Perfil Comercial vs R1.6 vs Diferencia"
        )

        grafico.y_axis.title = "Salidas"
        grafico.x_axis.title = (
            "Terminal / Servicio / Periodo"
        )

        # Series:
        # E = Salidas R1.6
        # F = Perfil Comercial
        # G = Diferencia

        datos_r16 = Reference(
            ws,
            min_col=5,
            min_row=3,
            max_row=ws.max_row,
        )

        datos_perfil = Reference(
            ws,
            min_col=6,
            min_row=3,
            max_row=ws.max_row,
        )

        datos_diferencia = Reference(
            ws,
            min_col=7,
            min_row=3,
            max_row=ws.max_row,
        )

        categorias = Reference(
            ws,
            min_col=2,
            min_row=4,
            max_row=ws.max_row,
        )

        grafico.add_data(
            datos_perfil,
            titles_from_data=True,
        )

        grafico.add_data(
            datos_r16,
            titles_from_data=True,
        )

        grafico.add_data(
            datos_diferencia,
            titles_from_data=True,
        )

        grafico.set_categories(
            categorias
        )

        grafico.height = 12
        grafico.width = 28

        ws.add_chart(
            grafico,
            "I3",
        )

    # ============================================================
    # DETALLE R1.6
    # ============================================================

    ws = wb.create_sheet(
        "DETALLE R1.6"
    )

    columnas_detalle = [
        "Fecha",
        "Unidad",
        "Terminal",
        "Servicio",
        "Código TS",
        "PPU",
        "Inicio servicio",
        "Fin servicio",
        "Sentido",
        "Ruta",
        "Periodo 30 min",
        "Periodo 1 hora",
    ]

    titulo(
        ws,
        "DETALLE DE SALIDAS COMERCIALES R1.6",
        len(columnas_detalle),
    )

    cabecera(
        ws,
        3,
        columnas_detalle,
    )

    for item in detalle:
        ws.append([
            item["fecha"],
            item["unidad"],
            item["terminal"],
            item["servicio"],
            item["codigo_ts"],
            item["ppu"],
            item["inicio_servicio"],
            item["fin_servicio"],
            item["sentido"],
            item["ruta"],
            item["periodo_30"],
            item["periodo_1h"],
        ])

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = ws.dimensions
    ajustar(ws, 32)

    # ============================================================
    # SALIDA
    # ============================================================

    buffer = BytesIO()

    wb.save(
        buffer
    )

    buffer.seek(0)

    nombre_archivo = (
        f"REPORTE_COMERCIAL_R16_"
        f"{unidad}_"
        f"{fecha}.xlsx"
    )

    return {
        "buffer": buffer,
        "nombre_archivo":
            nombre_archivo,
        "media_type":
            MIME_XLSX,
        "total_real":
            total_real,
        "total_perfil":
            total_perfil,
        "total_detalle":
            len(detalle),
        "terminales":
            len(resumen_terminal),
        "estado_cobertura":
            estado_cobertura,
    }
