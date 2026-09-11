from __future__ import annotations

from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.models import FlotaAsignacionTerminal


HOJAS_TERMINALES = {
    "EL RETIRO": "EL RETIRO",
    "STA. MARGARITA": "SANTA MARGARITA",
    "SANTA MARGARITA": "SANTA MARGARITA",
    "CONDELL": "CONDELL",
    "STA. MARTA": "SANTA MARTA",
    "SANTA MARTA": "SANTA MARTA",
    "A. LUCO": "AGUIRRE LUCO",
    "AGUIRRE LUCO": "AGUIRRE LUCO",
    "PIE ANDINO": "PIE ANDINO",
    "JUANITA": "JUANITA",
}

HOJAS_ESPECIALES = {
    "FLOTA SOPORTE": {
        "es_soporte": True,
        "es_auxiliar": False,
    },
    "FLOTA AUXILIAR": {
        "es_soporte": False,
        "es_auxiliar": True,
    },
}


def _texto(valor: Any) -> str:
    if valor is None:
        return ""

    return str(valor).strip()


def _normalizar(valor: Any) -> str:
    return _texto(valor).upper()


def normalizar_ppu(valor: Any) -> str:
    return (
        _normalizar(valor)
        .replace(" ", "")
        .replace(".", "")
    )


def _buscar_columna(encabezados: list[str], candidatos: list[str]) -> int | None:

    normalizados = {
        _normalizar(nombre): idx
        for idx, nombre in enumerate(encabezados)
        if nombre
    }

    for candidato in candidatos:
        candidato_n = _normalizar(candidato)

        if candidato_n in normalizados:
            return normalizados[candidato_n]

    return None


def normalizar_terminal(valor: Any) -> str:

    terminal = _normalizar(valor)

    equivalencias = {
        "STA. MARGARITA": "SANTA MARGARITA",
        "STA MARGARITA": "SANTA MARGARITA",
        "SANTA MARGARITA": "SANTA MARGARITA",

        "STA. MARTA": "SANTA MARTA",
        "STA MARTA": "SANTA MARTA",
        "SANTA MARTA": "SANTA MARTA",

        "A. LUCO": "AGUIRRE LUCO",
        "A LUCO": "AGUIRRE LUCO",
        "AGUIRRE LUCO": "AGUIRRE LUCO",

        "EL RETIRO": "EL RETIRO",
        "CONDELL": "CONDELL",
        "PIE ANDINO": "PIE ANDINO",
        "JUANITA": "JUANITA",
    }

    return equivalencias.get(
        terminal,
        terminal,
    )


def _determinar_unidad(terminal: str) -> str | None:

    terminal = normalizar_terminal(
        terminal
    )

    u8 = {
        "CONDELL",
        "EL RETIRO",
        "SANTA MARGARITA",
        "SANTA MARTA",
    }

    u9 = {
        "AGUIRRE LUCO",
        "JUANITA",
        "PIE ANDINO",
    }

    if terminal in u8:
        return "U8"

    if terminal in u9:
        return "U9"

    return None


def _ultima_version(db: Session) -> str | None:

    return (
        db.query(FlotaAsignacionTerminal.version)
        .order_by(
            FlotaAsignacionTerminal.version.desc()
        )
        .first()
        or [None]
    )[0]


def obtener_catalogo_actual(
    db: Session,
) -> dict[str, FlotaAsignacionTerminal]:

    version = _ultima_version(db)

    if not version:
        return {}

    filas = (
        db.query(FlotaAsignacionTerminal)
        .filter(
            FlotaAsignacionTerminal.version == version,
            FlotaAsignacionTerminal.activo.is_(True),
        )
        .all()
    )

    return {
        normalizar_ppu(x.ppu): x
        for x in filas
        if normalizar_ppu(x.ppu)
    }


def leer_excel_distribucion(
    ruta: str | Path,
) -> dict:

    ruta = Path(ruta)

    wb = load_workbook(
        ruta,
        read_only=True,
        data_only=True,
    )

    filas_finales: dict[str, dict] = {}

    duplicados = []
    invalidas = []

    # La Flota Auxiliar se reconoce desde el Excel,
    # pero NO pertenece al universo operacional utilizado
    # por Flota Operativa / Sin transmisi?n.
    auxiliares_fuera_universo = []

    hojas_leidas = []

    try:

        for nombre_hoja in wb.sheetnames:

            nombre_norm = _normalizar(nombre_hoja)

            terminal_fijo = HOJAS_TERMINALES.get(
                nombre_norm
            )

            config_especial = HOJAS_ESPECIALES.get(
                nombre_norm
            )

            if (
                not terminal_fijo
                and not config_especial
            ):
                continue

            ws = wb[nombre_hoja]

            hojas_leidas.append(nombre_hoja)

            filas = ws.iter_rows(
                values_only=True
            )

            try:
                encabezado_original = next(filas)
            except StopIteration:
                continue

            encabezados = [
                _texto(x)
                for x in encabezado_original
            ]

            idx_ppu = _buscar_columna(
                encabezados,
                [
                    "PPU",
                    "PLACA",
                    "PATENTE",
                    "PATENTE BUS",
                ],
            )

            idx_empresa = _buscar_columna(
                encabezados,
                [
                    "EMPRESA",
                ],
            )

            idx_interno = _buscar_columna(
                encabezados,
                [
                    "INTERNO",
                    "NRO INTERNO",
                    "NUMERO INTERNO",
                ],
            )

            idx_tipo_bus = _buscar_columna(
                encabezados,
                [
                    "TIPO BUS",
                    "TIPO DE BUS",
                ],
            )

            idx_tipo_flota = _buscar_columna(
                encabezados,
                [
                    "TIPO FLOTA",
                    "TIPO DE FLOTA",
                ],
            )

            idx_terminal = _buscar_columna(
                encabezados,
                [
                    "TERMINAL OP",
                    "TERMINAL APOYO",
                    "TERMINAL",
                    "TERMINAL OPERACION",
                ],
            )

            if idx_ppu is None:

                invalidas.append({
                    "hoja": nombre_hoja,
                    "error": "SIN COLUMNA PPU/PLACA",
                })

                continue

            for numero_fila, fila in enumerate(
                filas,
                start=2,
            ):

                valores = list(fila)

                def valor_indice(idx):
                    if idx is None:
                        return ""

                    if idx >= len(valores):
                        return ""

                    return valores[idx]

                ppu = normalizar_ppu(
                    valor_indice(idx_ppu)
                )

                if not ppu:
                    continue

                terminal = terminal_fijo

                if idx_terminal is not None:
                    terminal_excel = _normalizar(
                        valor_indice(idx_terminal)
                    )

                    if terminal_excel:
                        terminal = normalizar_terminal(
                            terminal_excel
                        )

                terminal = normalizar_terminal(
                    terminal or ""
                )

                es_soporte = False
                es_auxiliar = False

                if config_especial:
                    es_soporte = bool(
                        config_especial[
                            "es_soporte"
                        ]
                    )
                    es_auxiliar = bool(
                        config_especial[
                            "es_auxiliar"
                        ]
                    )

                unidad = _determinar_unidad(
                    terminal or ""
                )

                empresa = _normalizar(
                    valor_indice(
                        idx_empresa
                    )
                )

                registro = {
                    "ppu": ppu,
                    "terminal": terminal or "",
                    "unidad": unidad,
                    "empresa": empresa,
                    "interno": _texto(
                        valor_indice(
                            idx_interno
                        )
                    ),
                    "tipo_bus": _normalizar(
                        valor_indice(
                            idx_tipo_bus
                        )
                    ),
                    "tipo_flota": _normalizar(
                        valor_indice(
                            idx_tipo_flota
                        )
                    ),
                    "es_soporte": es_soporte,
                    "es_auxiliar": es_auxiliar,
                    "hoja": nombre_hoja,
                    "fila": numero_fila,
                }

                if not registro["terminal"]:
                    invalidas.append({
                        "hoja": nombre_hoja,
                        "fila": numero_fila,
                        "ppu": ppu,
                        "error": "SIN TERMINAL",
                    })

                    continue

                # --------------------------------------------------
                # FLOTA AUXILIAR
                #
                # Se conserva como informaci?n de diagn?stico,
                # pero queda fuera del universo operacional.
                # No alimenta Flota asignada ni Sin transmisi?n.
                # --------------------------------------------------
                if es_auxiliar:

                    auxiliares_fuera_universo.append(
                        registro
                    )

                    continue

                # --------------------------------------------------
                # FLOTA SOPORTE
                #
                # Una PPU de soporte ya existe previamente en su
                # hoja de terminal. No es un duplicado real:
                # la hoja Flota Soporte complementa / actualiza
                # el registro maestro y debe dejar es_soporte=True.
                # --------------------------------------------------
                if ppu in filas_finales:

                    anterior = filas_finales[ppu]

                    if es_soporte:

                        filas_finales[ppu] = registro

                        continue

                    duplicados.append({
                        "ppu": ppu,
                        "hoja_anterior":
                            anterior.get("hoja"),
                        "fila_anterior":
                            anterior.get("fila"),
                        "hoja_nueva":
                            nombre_hoja,
                        "fila_nueva":
                            numero_fila,
                    })

                    continue

                filas_finales[ppu] = registro

    finally:

        wb.close()

    por_unidad = Counter(
        x["unidad"] or "SIN UNIDAD"
        for x in filas_finales.values()
    )

    por_terminal = Counter(
        (
            x["unidad"] or "SIN UNIDAD",
            x["terminal"] or "SIN TERMINAL",
        )
        for x in filas_finales.values()
    )

    return {
        "archivo": ruta.name,
        "total_ppu": len(filas_finales),
        "hojas_leidas": hojas_leidas,
        "duplicados": duplicados,
        "invalidas": invalidas,

        "auxiliares_fuera_universo":
            auxiliares_fuera_universo,

        "total_auxiliares_fuera_universo":
            len(auxiliares_fuera_universo),

        "por_unidad": dict(
            sorted(
                por_unidad.items()
            )
        ),
        "por_terminal": [
            {
                "unidad": unidad,
                "terminal": terminal,
                "cantidad": cantidad,
            }
            for (
                unidad,
                terminal
            ), cantidad in sorted(
                por_terminal.items()
            )
        ],
        "registros": filas_finales,
    }


def comparar_con_catalogo(
    db: Session,
    resultado_excel: dict,
) -> dict:

    actual = obtener_catalogo_actual(
        db
    )

    nuevo = resultado_excel[
        "registros"
    ]

    actuales_ppu = set(
        actual.keys()
    )

    nuevas_ppu = set(
        nuevo.keys()
    )

    altas = sorted(
        nuevas_ppu - actuales_ppu
    )

    bajas = sorted(
        actuales_ppu - nuevas_ppu
    )

    cambios = []

    campos = [
        "terminal",
        "unidad",
        "empresa",
        "interno",
        "tipo_bus",
        "tipo_flota",
    ]

    for ppu in sorted(
        actuales_ppu & nuevas_ppu
    ):

        viejo = actual[ppu]
        nuevo_reg = nuevo[ppu]

        diferencias = {}

        for campo in campos:

            valor_viejo = _normalizar(
                getattr(
                    viejo,
                    campo,
                    None,
                )
            )

            valor_nuevo = _normalizar(
                nuevo_reg.get(
                    campo
                )
            )

            if valor_viejo != valor_nuevo:

                diferencias[campo] = {
                    "antes":
                        valor_viejo,
                    "despues":
                        valor_nuevo,
                }

        viejo_soporte = bool(
            viejo.es_soporte
        )

        nuevo_soporte = bool(
            nuevo_reg.get(
                "es_soporte"
            )
        )

        if viejo_soporte != nuevo_soporte:

            diferencias[
                "es_soporte"
            ] = {
                "antes": viejo_soporte,
                "despues": nuevo_soporte,
            }

        viejo_auxiliar = bool(
            viejo.es_auxiliar
        )

        nuevo_auxiliar = bool(
            nuevo_reg.get(
                "es_auxiliar"
            )
        )

        if viejo_auxiliar != nuevo_auxiliar:

            diferencias[
                "es_auxiliar"
            ] = {
                "antes": viejo_auxiliar,
                "despues": nuevo_auxiliar,
            }

        if diferencias:

            cambios.append({
                "ppu": ppu,
                "cambios": diferencias,
            })

    return {
        "version_actual":
            _ultima_version(db),

        "archivo_nuevo":
            resultado_excel[
                "archivo"
            ],

        "total_actual":
            len(actual),

        "total_nuevo":
            len(nuevo),

        "altas":
            len(altas),

        "bajas":
            len(bajas),

        "modificados":
            len(cambios),

        "sin_cambio":
            len(
                actuales_ppu
                & nuevas_ppu
            ) - len(cambios),

        "duplicados_archivo":
            len(
                resultado_excel[
                    "duplicados"
                ]
            ),

        "filas_invalidas":
            len(
                resultado_excel[
                    "invalidas"
                ]
            ),

        "detalle_altas":
            altas,

        "detalle_bajas":
            bajas,

        "detalle_cambios":
            cambios,

        "por_unidad":
            resultado_excel[
                "por_unidad"
            ],

        "por_terminal":
            resultado_excel[
                "por_terminal"
            ],

        "hojas_leidas":
            resultado_excel[
                "hojas_leidas"
            ],

        "duplicados":
            resultado_excel[
                "duplicados"
            ],

        "invalidas":
            resultado_excel[
                "invalidas"
            ],
    }


def previsualizar_importacion(
    db: Session,
    ruta: str | Path,
) -> dict:

    datos = leer_excel_distribucion(
        ruta
    )

    comparacion = comparar_con_catalogo(
        db,
        datos,
    )

    return {
        "estado": "PREVISUALIZACION",
        "fecha_revision":
            datetime.now().isoformat(
                timespec="seconds"
            ),
        **comparacion,
    }
