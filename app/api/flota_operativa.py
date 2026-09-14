from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session

from app.database import get_db
from app.services.flota_operativa_service import (
    consultar_detalle,
    consultar_flota,
    obtener_filtros,
)

router = APIRouter(prefix="/api/flota-operativa", tags=["Flota Operativa"])


@router.get("/filtros")
def filtros(db: Session = Depends(get_db)):
    return obtener_filtros(db)


@router.get("/consulta")
def consulta(
    fecha_desde: str | None = Query(default=None),
    fecha_hasta: str | None = Query(default=None),
    unidad: str | None = Query(default=None),
    terminal: str | None = Query(default=None),
    servicio: str | None = Query(default=None),
    periodo_inicial: int = Query(default=1, ge=1, le=24),
    periodo_final: int = Query(default=24, ge=1, le=24),
    db: Session = Depends(get_db),
):
    try:
        return consultar_flota(
            db=db,
            fecha_desde=fecha_desde,
            fecha_hasta=fecha_hasta,
            unidad=unidad,
            terminal=terminal,
            servicio=servicio,
            periodo_inicial=periodo_inicial,
            periodo_final=periodo_final,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))


@router.get("/detalle")
def detalle(
    fecha: str = Query(...),
    periodo: int = Query(..., ge=1, le=24),
    terminal: str = Query(...),
    servicio: str | None = Query(default=None),
    unidad: str | None = Query(default=None),
    db: Session = Depends(get_db),
):
    try:
        filas = consultar_detalle(
            db=db,
            fecha=fecha,
            periodo=periodo,
            terminal=terminal,
            servicio=servicio,
            unidad=unidad,
        )
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc))
    return {"total": len(filas), "registros": filas}

@router.get("/exportar-excel")
def exportar_excel(
    fecha_desde: str | None = Query(default=None),
    fecha_hasta: str | None = Query(default=None),
    unidad: str | None = Query(default=None),
    terminal: str | None = Query(default=None),
    servicio: str | None = Query(default=None),
    periodo_inicial: int = Query(default=1, ge=1, le=24),
    periodo_final: int = Query(default=24, ge=1, le=24),
    db: Session = Depends(get_db),
):

    from io import BytesIO
    from collections import defaultdict

    from fastapi.responses import StreamingResponse

    from openpyxl import Workbook
    from openpyxl.chart import (
        BarChart,
        LineChart,
        DoughnutChart,
        Reference,
    )
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
        Border,
        Side,
    )
    from openpyxl.utils import get_column_letter


    # ================================================================
    # HELPERS
    # ================================================================

    def periodo_txt(p):
        try:
            return f"P{int(p):02d}"
        except Exception:
            return str(p or "")


    def filtro_txt(valor, defecto="Todos"):
        if valor in (None, "", "null"):
            return defecto
        return str(valor)


    def estilizar_titulo(ws, titulo, subtitulo=None):

        ws.merge_cells("A1:H1")
        ws["A1"] = titulo
        ws["A1"].font = Font(
            bold=True,
            size=18,
            color="FFFFFF",
        )
        ws["A1"].fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )
        ws["A1"].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        ws.row_dimensions[1].height = 28

        if subtitulo:
            ws.merge_cells("A2:H2")
            ws["A2"] = subtitulo
            ws["A2"].font = Font(
                italic=True,
                size=10,
                color="666666",
            )
            ws["A2"].alignment = Alignment(
                horizontal="center",
            )


    def preparar_tabla(ws):

        azul = PatternFill(
            "solid",
            fgColor="1F4E78",
        )

        borde = Border(
            bottom=Side(
                style="thin",
                color="D9E2F3",
            )
        )

        for cell in ws[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = azul
            cell.alignment = Alignment(
                horizontal="center",
            )

        for row in ws.iter_rows():

            for cell in row:

                cell.border = borde

        for col in range(
            1,
            ws.max_column + 1,
        ):

            ancho = 12

            for cell in ws[
                get_column_letter(col)
            ]:

                valor = str(
                    cell.value
                    if cell.value is not None
                    else ""
                )

                ancho = max(
                    ancho,
                    min(
                        len(valor) + 2,
                        32,
                    )
                )

            ws.column_dimensions[
                get_column_letter(col)
            ].width = ancho


    def crear_dashboard(
        wb,
        nombre_hoja,
        titulo,
        datos,
        empresa=None,
    ):

        total = (
            datos.get("total_global")
            or {}
        )

        resumen_terminal_global = (
            datos.get("resumen_terminal_global")
            or []
        )

        resumen_terminal = (
            datos.get("resumen_terminal")
            or []
        )

        cobertura = (
            datos.get("cobertura_fuente")
            or []
        )

        ws = wb.create_sheet(
            nombre_hoja
        )

        estilizar_titulo(
            ws,
            titulo,
            (
                f"Per\u00edodo {periodo_txt(periodo_inicial)} "
                f"a {periodo_txt(periodo_final)}"
            ),
        )

        # ------------------------------------------------------------
        # FILTROS
        # ------------------------------------------------------------

        ws["A4"] = "Fecha desde"
        ws["B4"] = filtro_txt(fecha_desde, "-")

        ws["A5"] = "Fecha hasta"
        ws["B5"] = filtro_txt(fecha_hasta, "-")

        ws["A6"] = "Unidad"
        ws["B6"] = (
            empresa
            if empresa
            else filtro_txt(unidad, "Todas")
        )

        ws["A7"] = "Terminal"
        ws["B7"] = filtro_txt(
            terminal,
            "Todos",
        )

        ws["A8"] = "Servicio"
        ws["B8"] = filtro_txt(
            servicio,
            "Todos",
        )

        # ------------------------------------------------------------
        # KPIs
        # ------------------------------------------------------------

        kpis = [
            (
                "Flota asignada",
                int(
                    total.get("asignada")
                    or 0
                ),
            ),
            (
                "Buses en la calle",
                int(
                    total.get("operativa")
                    or 0
                ),
            ),
            (
                "Sin transmisi\u00f3n",
                int(
                    total.get("sin_transmision")
                    or 0
                ),
            ),
            (
                "% Operativa",
                float(
                    total.get("porcentaje_operativa")
                    or 0
                ) / 100,
            ),
        ]

        fila_kpi = 4

        for titulo_kpi, valor in kpis:

            ws.cell(
                fila_kpi,
                4,
                titulo_kpi,
            )

            ws.cell(
                fila_kpi,
                5,
                valor,
            )

            ws.cell(
                fila_kpi,
                4,
            ).font = Font(
                bold=True,
                color="FFFFFF",
            )

            ws.cell(
                fila_kpi,
                4,
            ).fill = PatternFill(
                "solid",
                fgColor="4472C4",
            )

            ws.cell(
                fila_kpi,
                5,
            ).font = Font(
                bold=True,
                size=14,
            )

            fila_kpi += 1

        ws["E7"].number_format = "0.00%"

        # ------------------------------------------------------------
        # COBERTURA / D?A PARCIAL
        # ------------------------------------------------------------

        parciales = [
            c
            for c in cobertura
            if not bool(
                c.get("dia_completo")
            )
        ]

        if parciales:

            ws.merge_cells(
                "A10:H10"
            )

            ws["A10"] = (
                "ADVERTENCIA: "
                "R1.6 DEL D?A ACTUAL ES PARCIAL"
            )

            ws["A10"].font = Font(
                bold=True,
                color="9C6500",
            )

            ws["A10"].fill = PatternFill(
                "solid",
                fgColor="FFEB9C",
            )

            ws["A10"].alignment = Alignment(
                horizontal="center",
            )

        # ------------------------------------------------------------
        # TABLA TERMINALES
        # ------------------------------------------------------------

        fila_tabla = 13

        encabezados = [
            "Terminal",
            "Asignada",
            "En calle",
            "Sin Tx",
            "% Operativa",
        ]

        for col, h in enumerate(
            encabezados,
            start=1,
        ):
            c = ws.cell(
                fila_tabla,
                col,
                h,
            )

            c.font = Font(
                bold=True,
                color="FFFFFF",
            )

            c.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )

        fila = fila_tabla + 1

        for r in resumen_terminal_global:

            ws.cell(
                fila,
                1,
                r.get("terminal_nombre")
                or r.get("terminal")
                or "",
            )

            ws.cell(
                fila,
                2,
                int(
                    r.get("asignada")
                    or 0
                ),
            )

            ws.cell(
                fila,
                3,
                int(
                    r.get("operativa")
                    or 0
                ),
            )

            ws.cell(
                fila,
                4,
                int(
                    r.get("sin_transmision")
                    or 0
                ),
            )

            ws.cell(
                fila,
                5,
                float(
                    r.get("porcentaje_operativa")
                    or 0
                ) / 100,
            )

            ws.cell(
                fila,
                5,
            ).number_format = "0.00%"

            fila += 1

        # ------------------------------------------------------------
        # GR?FICO BARRAS
        # ------------------------------------------------------------

        if fila > fila_tabla + 1:

            chart = BarChart()

            chart.type = "col"
            chart.style = 10

            chart.title = (
                "Flota asignada vs "
                "Buses en la calle por terminal"
            )

            chart.y_axis.title = "Buses"
            chart.x_axis.title = "Terminal"

            chart.height = 8
            chart.width = 16

            chart.legend.position = "b"

            datos_chart = Reference(
                ws,
                min_col=2,
                max_col=3,
                min_row=fila_tabla,
                max_row=fila - 1,
            )

            categorias = Reference(
                ws,
                min_col=1,
                min_row=fila_tabla + 1,
                max_row=fila - 1,
            )

            chart.add_data(
                datos_chart,
                titles_from_data=True,
            )

            chart.set_categories(
                categorias
            )

            ws.add_chart(
                chart,
                "G13",
            )

        # ------------------------------------------------------------
        # DONA ESTADO GENERAL
        # ------------------------------------------------------------

        ws["N4"] = "Estado"
        ws["O4"] = "Cantidad"

        ws["N5"] = "Buses en la calle"
        ws["O5"] = int(
            total.get("operativa")
            or 0
        )

        ws["N6"] = "Sin transmisi\u00f3n"
        ws["O6"] = int(
            total.get("sin_transmision")
            or 0
        )

        dona = DoughnutChart()

        dona.title = (
            "Estado general de la flota"
        )

        datos_dona = Reference(
            ws,
            min_col=15,
            min_row=4,
            max_row=6,
        )

        categorias_dona = Reference(
            ws,
            min_col=14,
            min_row=5,
            max_row=6,
        )

        dona.add_data(
            datos_dona,
            titles_from_data=True,
        )

        dona.set_categories(
            categorias_dona
        )

        dona.holeSize = 62
        dona.firstSliceAng = 270
        dona.height = 8
        dona.width = 10
        dona.legend.position = "b"

        ws.add_chart(
            dona,
            "W13",
        )

        # ------------------------------------------------------------
        # COMPORTAMIENTO P01-P24
        # ------------------------------------------------------------

        comportamiento = defaultdict(int)

        for r in resumen_terminal:

            p = r.get("periodo")

            if p is None:
                continue

            comportamiento[
                int(p)
            ] += int(
                r.get("operativa")
                or 0
            )

        fila_periodo = 32

        ws.cell(
            fila_periodo,
            1,
            "Per\u00edodo",
        )

        ws.cell(
            fila_periodo,
            2,
            "Flota operativa",
        )

        for c in (
            ws.cell(fila_periodo, 1),
            ws.cell(fila_periodo, 2),
        ):
            c.font = Font(
                bold=True,
                color="FFFFFF",
            )

            c.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )

        fila_p = fila_periodo + 1

        for p in range(
            int(periodo_inicial),
            int(periodo_final) + 1,
        ):

            ws.cell(
                fila_p,
                1,
                periodo_txt(p),
            )

            ws.cell(
                fila_p,
                2,
                comportamiento.get(
                    p,
                    0,
                ),
            )

            fila_p += 1

        linea = LineChart()

        linea.title = (
            "Comportamiento de la flota P01-P24"
        )

        linea.style = 13
        linea.height = 8
        linea.width = 18

        linea.y_axis.title = (
            "Flota operativa"
        )

        linea.x_axis.title = (
            "Per\u00edodo"
        )

        linea.legend.position = "b"

        data = Reference(
            ws,
            min_col=2,
            min_row=fila_periodo,
            max_row=fila_p - 1,
        )

        cats = Reference(
            ws,
            min_col=1,
            min_row=fila_periodo + 1,
            max_row=fila_p - 1,
        )

        linea.add_data(
            data,
            titles_from_data=True,
        )

        linea.set_categories(
            cats
        )

        for serie in linea.series:

            serie.marker.symbol = "circle"
            serie.marker.size = 6

        ws.add_chart(
            linea,
            "G32",
        )

        # ------------------------------------------------------------
        # COLUMN WIDTHS
        # ------------------------------------------------------------

        for col, ancho in {
            "A": 22,
            "B": 16,
            "C": 14,
            "D": 18,
            "E": 16,
            "F": 4,
        }.items():

            ws.column_dimensions[
                col
            ].width = ancho

        ws.freeze_panes = "A13"

        return ws


    # ================================================================
    # CONSULTAS
    # ================================================================

    datos_general = consultar_flota(
        db=db,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        unidad=unidad,
        terminal=terminal,
        servicio=servicio,
        periodo_inicial=periodo_inicial,
        periodo_final=periodo_final,
    )

    datos_u8 = consultar_flota(
        db=db,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        unidad="U8",
        terminal=terminal,
        servicio=servicio,
        periodo_inicial=periodo_inicial,
        periodo_final=periodo_final,
    )

    datos_u9 = consultar_flota(
        db=db,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        unidad="U9",
        terminal=terminal,
        servicio=servicio,
        periodo_inicial=periodo_inicial,
        periodo_final=periodo_final,
    )


    # ================================================================
    # R10 - HISTORICO DIARIO REAL
    # ================================================================

    from datetime import datetime, timedelta

    historico = {
        "U8": [],
        "U9": [],
    }

    fecha_hist_desde = fecha_desde
    fecha_hist_hasta = fecha_hasta

    if fecha_hist_desde and fecha_hist_hasta:

        f_ini = datetime.strptime(
            fecha_hist_desde,
            "%Y-%m-%d",
        ).date()

        f_fin = datetime.strptime(
            fecha_hist_hasta,
            "%Y-%m-%d",
        ).date()

        fecha_cursor = f_ini

        while fecha_cursor <= f_fin:

            fecha_txt = fecha_cursor.strftime(
                "%Y-%m-%d"
            )

            for unidad_hist in ("U8", "U9"):

                datos_dia = consultar_flota(
                    db=db,
                    fecha_desde=fecha_txt,
                    fecha_hasta=fecha_txt,
                    unidad=unidad_hist,
                    terminal=terminal,
                    servicio=servicio,
                    periodo_inicial=periodo_inicial,
                    periodo_final=periodo_final,
                )

                historico[
                    unidad_hist
                ].append({
                    "fecha": fecha_txt,
                    "datos": datos_dia,
                })

            fecha_cursor += timedelta(days=1)


    # ================================================================
    # HELPERS R10
    # ================================================================

    EMPRESA_UNIDAD = {
        "U8": "ALFA",
        "U9": "OMEGA",
    }


    def nombre_terminal(r):

        return (
            r.get("terminal_nombre")
            or r.get("terminal")
            or ""
        )


    def serie_periodos_terminal(datos):

        salida = defaultdict(
            lambda: defaultdict(int)
        )

        for r in (
            datos.get("resumen_terminal")
            or []
        ):

            p = r.get("periodo")

            if p is None:
                continue

            terminal_r = nombre_terminal(r)

            salida[
                terminal_r
            ][int(p)] += int(
                r.get("operativa")
                or 0
            )

        return salida


    def crear_resumen_historico(
        wb,
        unidad_nombre,
    ):

        empresa = EMPRESA_UNIDAD[
            unidad_nombre
        ]

        ws = wb.create_sheet(
            f"Resumen {unidad_nombre}"
        )

        ws.append([
            "Fecha",
            "Unidad",
            "Empresa",
            "Flota asignada",
            "Buses en la calle",
            "Sin transmisi\u00f3n",
            "% Operativa",
            "Cobertura",
        ])

        for item in historico[
            unidad_nombre
        ]:

            fecha_r = item["fecha"]
            datos_r = item["datos"]

            total = (
                datos_r.get("total_global")
                or {}
            )

            cobertura = (
                datos_r.get(
                    "cobertura_fuente"
                )
                or []
            )

            estado_cobertura = "COMPLETO"

            for c in cobertura:

                if not bool(
                    c.get("dia_completo")
                ):
                    estado_cobertura = (
                        "PARCIAL HASTA "
                        + str(
                            c.get("hasta")
                            or ""
                        )
                    )

            ws.append([
                fecha_r,
                unidad_nombre,
                empresa,
                int(
                    total.get("asignada")
                    or 0
                ),
                int(
                    total.get("operativa")
                    or 0
                ),
                int(
                    total.get(
                        "sin_transmision"
                    )
                    or 0
                ),
                float(
                    total.get(
                        "porcentaje_operativa"
                    )
                    or 0
                ) / 100,
                estado_cobertura,
            ])

        preparar_tabla(ws)

        for fila in range(
            2,
            ws.max_row + 1,
        ):
            ws.cell(
                fila,
                7,
            ).number_format = "0.00%"

        ws.freeze_panes = "A2"

        return ws


    def crear_comportamiento(
        wb,
        unidad_nombre,
        datos_unidad,
    ):

        empresa = EMPRESA_UNIDAD[
            unidad_nombre
        ]

        ws = wb.create_sheet(
            f"Comportamiento {unidad_nombre}"
        )

        serie_terminal = (
            serie_periodos_terminal(
                datos_unidad
            )
        )

        terminales = sorted(
            serie_terminal.keys()
        )

        ws.append(
            ["Per\u00edodo"]
            + terminales
            + [f"TOTAL {unidad_nombre}"]
        )

        for p in range(
            int(periodo_inicial),
            int(periodo_final) + 1,
        ):

            valores = [
                serie_terminal[t].get(
                    p,
                    0,
                )
                for t in terminales
            ]

            ws.append(
                [periodo_txt(p)]
                + valores
                + [sum(valores)]
            )

        preparar_tabla(ws)

        ws.freeze_panes = "A2"

        if ws.max_row >= 2:

            chart = LineChart()

            chart.title = (
                "Comportamiento de la flota "
                f"{unidad_nombre} - {empresa} "
                "por per\u00edodo"
            )

            chart.style = 13
            chart.height = 12
            chart.width = 23

            chart.y_axis.title = (
                "Buses operativos"
            )

            chart.x_axis.title = (
                "Per\u00edodo"
            )

            data = Reference(
                ws,
                min_col=2,
                max_col=ws.max_column,
                min_row=1,
                max_row=ws.max_row,
            )

            cats = Reference(
                ws,
                min_col=1,
                min_row=2,
                max_row=ws.max_row,
            )

            chart.add_data(
                data,
                titles_from_data=True,
            )

            chart.set_categories(
                cats
            )

            chart.legend.position = "r"

            for serie in chart.series:
                serie.marker.symbol = (
                    "circle"
                )
                serie.marker.size = 5

            ws.add_chart(
                chart,
                "A29",
            )

        return ws


    def crear_sin_tx_historico(
        wb,
        unidad_nombre,
    ):

        empresa = EMPRESA_UNIDAD[
            unidad_nombre
        ]

        ws = wb.create_sheet(
            f"Sin Tx Hist\u00f3rico {unidad_nombre}"
        )

        terminales = []

        for item in historico[
            unidad_nombre
        ]:

            for r in (
                item["datos"].get(
                    "resumen_terminal_global"
                )
                or []
            ):

                t = nombre_terminal(r)

                if (
                    t
                    and t not in terminales
                ):
                    terminales.append(t)

        terminales.sort()

        # ------------------------------------------------------------
        # MATRIZ HISTORICA PARA GRAFICO
        # ------------------------------------------------------------

        ws.append(
            ["Fecha"]
            + terminales
            + [f"TOTAL {unidad_nombre}"]
        )

        fila_matriz_inicio = 2

        for item in historico[
            unidad_nombre
        ]:

            por_terminal = {}

            for r in (
                item["datos"].get(
                    "resumen_terminal_global"
                )
                or []
            ):

                por_terminal[
                    nombre_terminal(r)
                ] = int(
                    r.get(
                        "sin_transmision"
                    )
                    or 0
                )

            valores = [
                por_terminal.get(
                    t,
                    0,
                )
                for t in terminales
            ]

            ws.append(
                [item["fecha"]]
                + valores
                + [sum(valores)]
            )

        fila_matriz_fin = ws.max_row

        preparar_tabla(ws)

        ws.freeze_panes = "A2"

        # ------------------------------------------------------------
        # GRAFICO DE SERIES HISTORICAS
        # ------------------------------------------------------------

        if fila_matriz_fin >= 2:

            chart = LineChart()

            chart.title = (
                "PPU sin transmisi\u00f3n por d\u00eda "
                f"y terminal - {unidad_nombre} "
                f"{empresa}"
            )

            chart.style = 13
            chart.height = 13
            chart.width = 25

            chart.y_axis.title = (
                "PPU sin transmisi\u00f3n"
            )

            chart.x_axis.title = "Fecha"

            data = Reference(
                ws,
                min_col=2,
                max_col=ws.max_column,
                min_row=1,
                max_row=fila_matriz_fin,
            )

            cats = Reference(
                ws,
                min_col=1,
                min_row=2,
                max_row=fila_matriz_fin,
            )

            chart.add_data(
                data,
                titles_from_data=True,
            )

            chart.set_categories(
                cats
            )

            chart.legend.position = "r"

            for serie in chart.series:
                serie.marker.symbol = (
                    "circle"
                )
                serie.marker.size = 6

            ws.add_chart(
                chart,
                "A8",
            )


        # ------------------------------------------------------------
        # DETALLE PPU SIN TRANSMISION POR FECHA
        # ------------------------------------------------------------

        fila_detalle = max(
            fila_matriz_fin + 20,
            30,
        )

        headers_detalle = [
            "Fecha",
            "Unidad",
            "Empresa",
            "Terminal",
            "PPU",
            "Interno",
            "Tipo Bus",
            "Tipo Flota",
            "Estado",
        ]

        for col, valor in enumerate(
            headers_detalle,
            start=1,
        ):

            cell = ws.cell(
                fila_detalle,
                col,
                valor,
            )

            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )

            cell.alignment = Alignment(
                horizontal="center",
            )

        fila = fila_detalle + 1

        for item in historico[
            unidad_nombre
        ]:

            fecha_r = item["fecha"]

            for r in (
                item["datos"].get(
                    "sin_transmision_detalle"
                )
                or []
            ):

                ws.cell(
                    fila,
                    1,
                    fecha_r,
                )

                ws.cell(
                    fila,
                    2,
                    unidad_nombre,
                )

                ws.cell(
                    fila,
                    3,
                    empresa,
                )

                ws.cell(
                    fila,
                    4,
                    nombre_terminal(r),
                )

                ws.cell(
                    fila,
                    5,
                    r.get("ppu"),
                )

                ws.cell(
                    fila,
                    6,
                    r.get("interno"),
                )

                ws.cell(
                    fila,
                    7,
                    r.get("tipo_bus"),
                )

                ws.cell(
                    fila,
                    8,
                    r.get("tipo_flota"),
                )

                ws.cell(
                    fila,
                    9,
                    r.get("estado"),
                )

                fila += 1

        for col in range(
            1,
            10,
        ):

            ws.column_dimensions[
                get_column_letter(col)
            ].width = 20

        ws.column_dimensions["A"].width = 14
        ws.column_dimensions["D"].width = 22
        ws.column_dimensions["E"].width = 14

        return ws


    # ================================================================
    # R15D - REPORTE SIN TRANSMISION POR UNIDAD Y TERMINAL
    #
    # Fuente:
    # - resumen_sin_transmision_dias
    # - sin_transmision_detalle_dias
    #
    # No recalcula Sin Tx.
    # Reutiliza el resultado certificado de consultar_flota().
    # ================================================================

    def crear_resumen_sin_tx(
        wb,
        datos_u8,
        datos_u9,
    ):

        ws = wb.create_sheet(
            "Resumen Sin Tx"
        )

        encabezados = [
            "Unidad",
            "Empresa",
            "Terminal",
            "Flota asignada",
            "PPU sin transmisi?n",
            "% sin transmisi?n",
            "D?as rango",
            "D?as evaluables",
            "D?as sin datos",
        ]

        ws.append(encabezados)

        unidades_datos = [
            ("U8", "ALFA", datos_u8),
            ("U9", "OMEGA", datos_u9),
        ]

        for unidad_r, empresa_r, datos_r in unidades_datos:

            resumen = (
                datos_r.get(
                    "resumen_sin_transmision_dias"
                )
                or []
            )

            dias_rango = int(
                datos_r.get(
                    "dias_rango_sin_transmision"
                )
                or 0
            )

            dias_evaluables = int(
                datos_r.get(
                    "dias_evaluables_sin_transmision"
                )
                or 0
            )

            dias_sin_datos = int(
                datos_r.get(
                    "dias_sin_datos_sin_transmision"
                )
                or 0
            )

            for fila in resumen:

                asignada = int(
                    fila.get(
                        "asignada"
                    )
                    or fila.get(
                        "total_ppu"
                    )
                    or 0
                )

                sin_tx = int(
                    fila.get(
                        "sin_transmision"
                    )
                    or 0
                )

                porcentaje_sin_tx = (
                    float(
                        fila.get(
                            "porcentaje_sin_transmision"
                        )
                        or 0
                    )
                    /
                    100
                )

                ws.append([
                    unidad_r,
                    empresa_r,
                    fila.get(
                        "terminal_nombre"
                    )
                    or fila.get(
                        "terminal"
                    )
                    or "",
                    asignada,
                    sin_tx,
                    porcentaje_sin_tx,
                    dias_rango,
                    dias_evaluables,
                    dias_sin_datos,
                ])

        preparar_tabla(ws)

        for fila in range(
            2,
            ws.max_row + 1,
        ):
            ws.cell(
                fila,
                6,
            ).number_format = "0.00%"

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["D"].width = 16
        ws.column_dimensions["E"].width = 20
        ws.column_dimensions["F"].width = 20
        ws.column_dimensions["G"].width = 14
        ws.column_dimensions["H"].width = 16
        ws.column_dimensions["I"].width = 14

        return ws


    def crear_detalle_sin_tx(
        wb,
        unidad_r,
        empresa_r,
        datos_r,
    ):

        nombre = (
            f"Sin Tx {unidad_r} {empresa_r}"
        )

        ws = wb.create_sheet(
            nombre
        )

        encabezados = [
            "Unidad",
            "Empresa",
            "Terminal",
            "PPU",
            "Interno",
            "Tipo Bus",
            "Tipo Flota",
            "Plazas",
            "D?as rango",
            "D?as evaluables",
            "D?as con transmisi?n",
            "D?as sin transmisi?n",
            "D?as sin datos",
            "% d?as sin transmisi?n",
            "Estado",
        ]

        ws.append(encabezados)

        detalle = (
            datos_r.get(
                "sin_transmision_detalle_dias"
            )
            or []
        )

        for fila in detalle:

            dias_rango = int(
                fila.get(
                    "dias_rango"
                )
                or 0
            )

            dias_evaluables = int(
                fila.get(
                    "dias_evaluables"
                )
                or 0
            )

            dias_con_tx = int(
                fila.get(
                    "dias_con_transmision"
                )
                or 0
            )

            dias_sin_tx = int(
                fila.get(
                    "dias_sin_transmision"
                )
                or 0
            )

            dias_sin_datos = int(
                fila.get(
                    "dias_sin_datos"
                )
                or 0
            )

            porcentaje_dias_sin_tx = (
                (
                    dias_sin_tx
                    /
                    dias_evaluables
                )
                if dias_evaluables > 0
                else 0
            )

            ws.append([
                unidad_r,
                empresa_r,
                fila.get(
                    "terminal_nombre"
                )
                or fila.get(
                    "terminal"
                )
                or "",
                fila.get(
                    "ppu"
                )
                or "",
                fila.get(
                    "interno"
                )
                or "",
                fila.get(
                    "tipo_bus"
                )
                or "",
                fila.get(
                    "tipo_flota"
                )
                or "",
                fila.get(
                    "plazas"
                ),
                dias_rango,
                dias_evaluables,
                dias_con_tx,
                dias_sin_tx,
                dias_sin_datos,
                porcentaje_dias_sin_tx,
                fila.get(
                    "estado"
                )
                or "SIN TRANSMISI?N",
            ])

        preparar_tabla(ws)

        for fila in range(
            2,
            ws.max_row + 1,
        ):
            ws.cell(
                fila,
                14,
            ).number_format = "0.00%"

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        anchos = {
            "A": 12,
            "B": 14,
            "C": 24,
            "D": 14,
            "E": 12,
            "F": 16,
            "G": 16,
            "H": 10,
            "I": 12,
            "J": 16,
            "K": 20,
            "L": 20,
            "M": 14,
            "N": 22,
            "O": 20,
        }

        for col, ancho in anchos.items():
            ws.column_dimensions[
                col
            ].width = ancho

        return ws


    # ================================================================
    # WORKBOOK R10
    # ================================================================

    wb = Workbook()

    hoja_default = wb.active
    wb.remove(hoja_default)


    # ================================================================
    # 01 DASHBOARD U8 - ALFA
    # ================================================================

    crear_dashboard(
        wb,
        "Dashboard U8 - ALFA",
        "SWAV - FLOTA OPERATIVA - U8 ALFA",
        datos_u8,
        "U8 - ALFA",
    )


    # ================================================================
    # 02 RESUMEN U8
    # ================================================================

    crear_resumen_historico(
        wb,
        "U8",
    )


    # ================================================================
    # 03 COMPORTAMIENTO U8
    # ================================================================

    crear_comportamiento(
        wb,
        "U8",
        datos_u8,
    )


    # ================================================================
    # 04 SIN TX HISTORICO U8
    # ================================================================

    crear_sin_tx_historico(
        wb,
        "U8",
    )


    # ================================================================
    # 05 DASHBOARD U9 - OMEGA
    # ================================================================

    crear_dashboard(
        wb,
        "Dashboard U9 - OMEGA",
        "SWAV - FLOTA OPERATIVA - U9 OMEGA",
        datos_u9,
        "U9 - OMEGA",
    )


    # ================================================================
    # 06 RESUMEN U9
    # ================================================================

    crear_resumen_historico(
        wb,
        "U9",
    )


    # ================================================================
    # 07 COMPORTAMIENTO U9
    # ================================================================

    crear_comportamiento(
        wb,
        "U9",
        datos_u9,
    )


    # ================================================================
    # 08 SIN TX HISTORICO U9
    # ================================================================

    crear_sin_tx_historico(
        wb,
        "U9",
    )


    # ================================================================
    # R15D - RESUMEN SIN TRANSMISION
    # ================================================================

    crear_resumen_sin_tx(
        wb,
        datos_u8,
        datos_u9,
    )


    # ================================================================
    # R15D - DETALLE U8 ALFA
    # ================================================================

    crear_detalle_sin_tx(
        wb,
        "U8",
        "ALFA",
        datos_u8,
    )


    # ================================================================
    # R15D - DETALLE U9 OMEGA
    # ================================================================

    crear_detalle_sin_tx(
        wb,
        "U9",
        "OMEGA",
        datos_u9,
    )


    # ================================================================
    # COBERTURA R1.6
    # Se mantiene como hoja tecnica adicional.
    # ================================================================

    ws_c = wb.create_sheet(
        "Cobertura R1.6"
    )

    ws_c.append([
        "Unidad",
        "Fecha",
        "Hasta",
        "Archivo",
        "D\u00eda completo",
    ])

    cobertura_general = (
        datos_general.get(
            "cobertura_fuente"
        )
        or []
    )

    for r in cobertura_general:

        ws_c.append([
            r.get("unidad"),
            r.get("fecha"),
            r.get("hasta"),
            r.get("archivo"),
            r.get("dia_completo"),
        ])

    preparar_tabla(ws_c)


    # ================================================================
    # SALIDA
    # ================================================================

    buffer = BytesIO()

    wb.save(
        buffer
    )

    buffer.seek(0)

    desde_nombre = (
        fecha_desde
        or "inicio"
    )

    hasta_nombre = (
        fecha_hasta
        or "fin"
    )

    nombre_archivo = (
        f"SWAV_FLOTA_OPERATIVA_"
        f"{desde_nombre}_"
        f"{hasta_nombre}.xlsx"
    )

    headers = {
        "Content-Disposition":
            f'attachment; filename="{nombre_archivo}"'
    }

    return StreamingResponse(
        buffer,
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers=headers,
    )


# ====================================================================
# R15HB - EXPORTACION HISTORICA COMPLETA SIN TX
#
# - Lee todas las fechas realmente existentes en HistoricoFlotaOperativa.
# - Solo exporta PPU que tengan al menos un dia SIN TX.
# - Agrupa por Unidad -> Terminal -> PPU.
# - Los dias sin cobertura NO cuentan.
# - No modifica historico, R1.6 ni Velocidades.
# ====================================================================

@router.get("/exportar-sin-tx-historico")
def exportar_sin_tx_historico(
    db: Session = Depends(get_db),
):

    from collections import defaultdict
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter
    from sqlalchemy import func

    from app.models import HistoricoFlotaOperativa


    EMPRESAS = {
        "U8": "ALFA",
        "U9": "OMEGA",
    }


    # ================================================================
    # FECHAS REALES EXISTENTES EN LA BD
    # ================================================================

    fechas_por_unidad = {}

    for unidad_r in ("U8", "U9"):

        filas_fecha = (
            db.query(
                HistoricoFlotaOperativa.fecha
            )
            .filter(
                HistoricoFlotaOperativa.unidad
                ==
                unidad_r
            )
            .distinct()
            .order_by(
                HistoricoFlotaOperativa.fecha
            )
            .all()
        )

        fechas_por_unidad[unidad_r] = [
            fila[0]
            for fila in filas_fecha
            if fila[0] is not None
        ]


    todas_fechas = sorted({
        fecha
        for fechas in fechas_por_unidad.values()
        for fecha in fechas
    })

    if not todas_fechas:
        raise HTTPException(
            status_code=404,
            detail=(
                "No existen datos historicos "
                "de Flota Operativa."
            ),
        )


    # ================================================================
    # OBTENER SIN TX DIA POR DIA
    #
    # Reutiliza consultar_flota().
    # De esta manera usa exactamente la misma regla certificada
    # que alimenta la pantalla PPU Sin Tx por dia y terminal.
    # ================================================================

    eventos = {}

    for unidad_r in ("U8", "U9"):

        for fecha_r in fechas_por_unidad[
            unidad_r
        ]:

            fecha_txt = (
                fecha_r.isoformat()
                if hasattr(
                    fecha_r,
                    "isoformat"
                )
                else str(fecha_r)
            )

            datos_dia = consultar_flota(
                db=db,
                fecha_desde=fecha_txt,
                fecha_hasta=fecha_txt,
                unidad=unidad_r,
                terminal=None,
                servicio=None,
                periodo_inicial=1,
                periodo_final=24,
            )

            detalle = (
                datos_dia.get(
                    "sin_transmision_detalle_dias"
                )
                or []
            )

            for fila in detalle:

                ppu = str(
                    fila.get("ppu")
                    or ""
                ).strip().upper()

                if not ppu:
                    continue

                terminal = str(
                    fila.get(
                        "terminal_nombre"
                    )
                    or
                    fila.get("terminal")
                    or ""
                ).strip()

                clave = (
                    unidad_r,
                    fecha_txt,
                    terminal.upper(),
                    ppu,
                )

                eventos[clave] = {
                    "fecha": fecha_txt,
                    "unidad": unidad_r,
                    "empresa":
                        EMPRESAS.get(
                            unidad_r,
                            unidad_r,
                        ),
                    "terminal": terminal,
                    "ppu": ppu,
                    "interno":
                        fila.get("interno")
                        or "",
                    "tipo_bus":
                        fila.get("tipo_bus")
                        or "",
                    "tipo_flota":
                        fila.get("tipo_flota")
                        or "",
                    "plazas":
                        fila.get("plazas"),
                }


    eventos = list(
        eventos.values()
    )


    # ================================================================
    # AGRUPAR POR PPU
    # ================================================================

    agrupado = {}

    for evento in eventos:

        clave = (
            evento["unidad"],
            evento["terminal"].upper(),
            evento["ppu"],
        )

        if clave not in agrupado:

            agrupado[clave] = {
                "unidad":
                    evento["unidad"],
                "empresa":
                    evento["empresa"],
                "terminal":
                    evento["terminal"],
                "ppu":
                    evento["ppu"],
                "interno":
                    evento["interno"],
                "tipo_bus":
                    evento["tipo_bus"],
                "tipo_flota":
                    evento["tipo_flota"],
                "plazas":
                    evento["plazas"],
                "fechas":
                    set(),
            }

        item = agrupado[clave]

        item["fechas"].add(
            evento["fecha"]
        )

        # Completar datos si en alguna fila venian vacios.

        if not item["interno"]:
            item["interno"] = (
                evento["interno"]
            )

        if not item["tipo_bus"]:
            item["tipo_bus"] = (
                evento["tipo_bus"]
            )

        if not item["tipo_flota"]:
            item["tipo_flota"] = (
                evento["tipo_flota"]
            )

        if item["plazas"] in (
            None,
            "",
        ):
            item["plazas"] = (
                evento["plazas"]
            )


    # ================================================================
    # INDICE DE DIAS EVALUABLES
    #
    # La racha se calcula sobre los dias disponibles en la BD.
    # Ejemplo:
    # 09,10,11,13,14
    #
    # El 12 no rompe la racha porque no existe como dia evaluable.
    # ================================================================

    indice_fecha_unidad = {}

    for unidad_r, fechas in (
        fechas_por_unidad.items()
    ):

        ordenadas = sorted(
            fechas
        )

        indice_fecha_unidad[
            unidad_r
        ] = {
            (
                fecha.isoformat()
                if hasattr(
                    fecha,
                    "isoformat"
                )
                else str(fecha)
            ): indice
            for indice, fecha
            in enumerate(ordenadas)
        }


    def calcular_racha(
        unidad_r,
        fechas_sin_tx,
    ):

        mapa = (
            indice_fecha_unidad.get(
                unidad_r,
                {}
            )
        )

        indices = sorted({
            mapa[fecha]
            for fecha in fechas_sin_tx
            if fecha in mapa
        })

        if not indices:
            return 0

        maxima = 1
        actual = 1

        for anterior, siguiente in zip(
            indices,
            indices[1:],
        ):

            if siguiente == anterior + 1:
                actual += 1
            else:
                actual = 1

            maxima = max(
                maxima,
                actual,
            )

        return maxima


    ranking = []

    for item in agrupado.values():

        fechas = sorted(
            item["fechas"]
        )

        dias_sin_tx = len(
            fechas
        )

        total_evaluables = len(
            fechas_por_unidad.get(
                item["unidad"],
                [],
            )
        )

        porcentaje = (
            (
                dias_sin_tx
                /
                total_evaluables
            )
            if total_evaluables > 0
            else 0
        )

        ranking.append({
            **item,
            "dias_evaluables":
                total_evaluables,
            "dias_sin_tx":
                dias_sin_tx,
            "porcentaje_sin_tx":
                porcentaje,
            "racha_maxima":
                calcular_racha(
                    item["unidad"],
                    fechas,
                ),
            "primera_fecha":
                fechas[0]
                if fechas
                else "",
            "ultima_fecha":
                fechas[-1]
                if fechas
                else "",
            "fechas_txt":
                ", ".join(
                    fechas
                ),
        })


    # Ranking general:
    # primero mas dias Sin Tx,
    # luego mayor racha,
    # luego unidad/terminal/PPU.

    ranking.sort(
        key=lambda x: (
            -x["dias_sin_tx"],
            -x["racha_maxima"],
            x["unidad"],
            x["terminal"],
            x["ppu"],
        )
    )


    # ================================================================
    # RESUMEN POR UNIDAD / TERMINAL
    # ================================================================

    resumen_terminal = defaultdict(
        lambda: {
            "ppus": set(),
            "eventos": 0,
            "max_dias": 0,
        }
    )

    for item in ranking:

        clave = (
            item["unidad"],
            item["empresa"],
            item["terminal"],
        )

        resumen_terminal[
            clave
        ]["ppus"].add(
            item["ppu"]
        )

        resumen_terminal[
            clave
        ]["eventos"] += (
            item["dias_sin_tx"]
        )

        resumen_terminal[
            clave
        ]["max_dias"] = max(
            resumen_terminal[
                clave
            ]["max_dias"],
            item["dias_sin_tx"],
        )


    # ================================================================
    # R15H-G - EXCEL HISTORICO SIN TX PROFESIONAL
    # ================================================================

    wb = Workbook()

    ws_default = wb.active
    wb.remove(ws_default)

    # ---------------------------------------------------------------
    # COLORES
    # ---------------------------------------------------------------

    COLOR_AZUL_OSCURO = "17365D"
    COLOR_AZUL = "1F4E78"
    COLOR_AZUL_MEDIO = "5B9BD5"
    COLOR_AZUL_CLARO = "D9EAF7"

    COLOR_BLANCO = "FFFFFF"
    COLOR_TEXTO = "1F1F1F"
    COLOR_GRIS = "666666"
    COLOR_GRIS_CLARO = "F2F2F2"

    COLOR_ROJO = "F4CCCC"
    COLOR_ROJO_FUERTE = "C00000"

    COLOR_AMARILLO = "FFF2CC"
    COLOR_AMARILLO_FUERTE = "BF9000"

    COLOR_VERDE = "D9EAD3"
    COLOR_VERDE_FUERTE = "38761D"

    COLOR_BORDE = "B4C6E7"

    borde_fino = Border(
        left=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
        right=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
        top=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
        bottom=Side(
            style="thin",
            color=COLOR_BORDE,
        ),
    )

    # ---------------------------------------------------------------
    # TEXTOS UNICODE SEGUROS
    #
    # Se usan escapes para evitar que consola/editor vuelva a
    # convertir vocales acentuadas en signos de interrogacion.
    # ---------------------------------------------------------------

    TXT_DIAS = "D\u00edas"
    TXT_MAXIMO = "M\u00e1ximo"
    TXT_MAXIMA = "M\u00e1xima"
    TXT_ULTIMA = "\u00daltima"
    TXT_HISTORICO = "Hist\u00f3rico"
    TXT_TRANSMISION = "Transmisi\u00f3n"

    # ---------------------------------------------------------------
    # UTILIDADES
    # ---------------------------------------------------------------

    def aplicar_encabezado(
        ws,
        fila,
        desde_col,
        hasta_col,
    ):

        for col in range(
            desde_col,
            hasta_col + 1,
        ):

            celda = ws.cell(
                fila,
                col,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_BLANCO,
                size=10,
            )

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_AZUL,
            )

            celda.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

            celda.border = borde_fino

        ws.row_dimensions[fila].height = 32


    def aplicar_tabla(
        ws,
        fila_inicio,
        fila_fin,
        columna_inicio,
        columna_fin,
    ):

        for fila in range(
            fila_inicio,
            fila_fin + 1,
        ):

            for col in range(
                columna_inicio,
                columna_fin + 1,
            ):

                celda = ws.cell(
                    fila,
                    col,
                )

                celda.border = borde_fino

                celda.alignment = Alignment(
                    vertical="center",
                )

                if fila % 2 == 0:

                    celda.fill = PatternFill(
                        "solid",
                        fgColor="F8FBFE",
                    )


    def ajustar_columnas(
        ws,
        maximo=38,
    ):

        for col in range(
            1,
            ws.max_column + 1,
        ):

            letra = get_column_letter(
                col
            )

            ancho = 10

            for celda in ws[
                letra
            ]:

                valor = (
                    ""
                    if celda.value is None
                    else str(celda.value)
                )

                ancho = max(
                    ancho,
                    min(
                        len(valor) + 2,
                        maximo,
                    ),
                )

            ws.column_dimensions[
                letra
            ].width = ancho


    def semaforo_porcentaje(
        celda,
        porcentaje,
    ):

        if porcentaje >= 0.60:

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_ROJO,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_ROJO_FUERTE,
            )

        elif porcentaje >= 0.30:

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_AMARILLO,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_AMARILLO_FUERTE,
            )

        else:

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_VERDE,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_VERDE_FUERTE,
            )


    def semaforo_dias(
        celda,
        dias,
    ):

        if dias >= 4:

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_ROJO,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_ROJO_FUERTE,
            )

        elif dias >= 2:

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_AMARILLO,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_AMARILLO_FUERTE,
            )

        else:

            celda.fill = PatternFill(
                "solid",
                fgColor=COLOR_VERDE,
            )

            celda.font = Font(
                bold=True,
                color=COLOR_VERDE_FUERTE,
            )


    # ---------------------------------------------------------------
    # RESUMENES PARA PRESENTACION
    # ---------------------------------------------------------------

    total_eventos = len(
        eventos
    )

    total_ppu_sin_tx = len({
        item["ppu"]
        for item in ranking
    })

    total_dias_historicos = len(
        todas_fechas
    )

    ppus_u8 = len({
        item["ppu"]
        for item in ranking
        if item["unidad"] == "U8"
    })

    ppus_u9 = len({
        item["ppu"]
        for item in ranking
        if item["unidad"] == "U9"
    })

    # ---------------------------------------------------------------
    # 1. RESUMEN EJECUTIVO
    # ---------------------------------------------------------------

    ws = wb.create_sheet(
        "Resumen Ejecutivo"
    )

    ws.sheet_view.showGridLines = False

    ws.merge_cells(
        "A1:H2"
    )

    titulo = ws["A1"]

    titulo.value = (
        "SWAV - REPORTE "
        f"{TXT_HISTORICO.upper()} "
        "PPU SIN TX"
    )

    titulo.font = Font(
        bold=True,
        color=COLOR_BLANCO,
        size=20,
    )

    titulo.fill = PatternFill(
        "solid",
        fgColor=COLOR_AZUL_OSCURO,
    )

    titulo.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 28

    primera_txt = (
        todas_fechas[0].isoformat()
        if hasattr(
            todas_fechas[0],
            "isoformat",
        )
        else str(
            todas_fechas[0]
        )
    )

    ultima_txt = (
        todas_fechas[-1].isoformat()
        if hasattr(
            todas_fechas[-1],
            "isoformat",
        )
        else str(
            todas_fechas[-1]
        )
    )

    ws.merge_cells(
        "A3:H3"
    )

    ws["A3"] = (
        f"Per\u00edodo disponible en BD: "
        f"{primera_txt} al {ultima_txt}  |  "
        f"{total_dias_historicos} d\u00edas evaluables"
    )

    ws["A3"].alignment = Alignment(
        horizontal="center",
    )

    ws["A3"].font = Font(
        italic=True,
        color=COLOR_GRIS,
    )

    tarjetas = [
        (
            "A5:B5",
            "A6:B7",
            "PPU CON SIN TX",
            total_ppu_sin_tx,
        ),
        (
            "C5:D5",
            "C6:D7",
            "EVENTOS PPU-D\u00cdA",
            total_eventos,
        ),
        (
            "E5:F5",
            "E6:F7",
            f"{TXT_DIAS.upper()} EVALUABLES",
            total_dias_historicos,
        ),
        (
            "G5:H5",
            "G6:H7",
            "UNIDADES",
            2,
        ),
    ]

    for (
        rango_titulo,
        rango_valor,
        etiqueta,
        valor,
    ) in tarjetas:

        ws.merge_cells(
            rango_titulo
        )

        ws.merge_cells(
            rango_valor
        )

        c_titulo = ws[
            rango_titulo.split(":")[0]
        ]

        c_valor = ws[
            rango_valor.split(":")[0]
        ]

        c_titulo.value = etiqueta

        c_titulo.fill = PatternFill(
            "solid",
            fgColor=COLOR_AZUL,
        )

        c_titulo.font = Font(
            bold=True,
            color=COLOR_BLANCO,
        )

        c_titulo.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        c_valor.value = valor

        c_valor.fill = PatternFill(
            "solid",
            fgColor=COLOR_AZUL_CLARO,
        )

        c_valor.font = Font(
            bold=True,
            size=22,
            color=COLOR_AZUL_OSCURO,
        )

        c_valor.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    ws["A9"] = "Unidad"
    ws["B9"] = "Empresa"
    ws["C9"] = "PPU con Sin TX"
    ws["D9"] = "Eventos PPU-D\u00eda"

    aplicar_encabezado(
        ws,
        9,
        1,
        4,
    )

    for unidad_r in (
        "U8",
        "U9",
    ):

        filas_unidad = [
            item
            for item in ranking
            if item["unidad"] == unidad_r
        ]

        eventos_unidad = [
            evento
            for evento in eventos
            if evento["unidad"] == unidad_r
        ]

        ws.append([
            unidad_r,
            EMPRESAS.get(
                unidad_r,
                unidad_r,
            ),
            len({
                item["ppu"]
                for item in filas_unidad
            }),
            len(
                eventos_unidad
            ),
        ])

    aplicar_tabla(
        ws,
        10,
        ws.max_row,
        1,
        4,
    )

    ws["A14"] = (
        "Nota: SIN TX significa que la PPU no registr\u00f3 "
        "transmisi\u00f3n en un d\u00eda evaluable. "
        "No significa autom\u00e1ticamente que el bus estuvo en taller."
    )

    ws.merge_cells(
        "A14:H15"
    )

    ws["A14"].alignment = Alignment(
        wrap_text=True,
        vertical="center",
    )

    ws["A14"].font = Font(
        italic=True,
        color=COLOR_GRIS,
    )

    ws["A14"].fill = PatternFill(
        "solid",
        fgColor=COLOR_GRIS_CLARO,
    )

    for col in range(
        1,
        9,
    ):
        ws.column_dimensions[
            get_column_letter(col)
        ].width = 18

    # ---------------------------------------------------------------
    # 2. RESUMEN TERMINALES
    # ---------------------------------------------------------------

    ws = wb.create_sheet(
        "Resumen Terminales"
    )

    encabezados = [
        "Unidad",
        "Empresa",
        "Terminal",
        "PPU con al menos 1 d\u00eda Sin TX",
        "Total PPU-D\u00edas Sin TX",
        "M\u00e1ximo d\u00edas Sin TX de un PPU",
        "Promedio d\u00edas Sin TX por PPU",
        "D\u00edas evaluables BD",
        "% promedio Sin TX",
        "Estado",
    ]

    ws.append(
        encabezados
    )

    for clave in sorted(
        resumen_terminal.keys()
    ):

        unidad_r, empresa_r, terminal_r = clave

        info = resumen_terminal[
            clave
        ]

        cantidad_ppu = len(
            info["ppus"]
        )

        eventos_terminal = (
            info["eventos"]
        )

        promedio = (
            eventos_terminal
            /
            cantidad_ppu
            if cantidad_ppu
            else 0
        )

        dias_evaluables = len(
            fechas_por_unidad.get(
                unidad_r,
                [],
            )
        )

        porcentaje_promedio = (
            promedio
            /
            dias_evaluables
            if dias_evaluables
            else 0
        )

        if porcentaje_promedio >= 0.60:
            estado = "ALTO"
        elif porcentaje_promedio >= 0.30:
            estado = "ATENCI\u00d3N"
        else:
            estado = "BAJO"

        ws.append([
            unidad_r,
            empresa_r,
            terminal_r,
            cantidad_ppu,
            eventos_terminal,
            info["max_dias"],
            promedio,
            dias_evaluables,
            porcentaje_promedio,
            estado,
        ])

    aplicar_encabezado(
        ws,
        1,
        1,
        ws.max_column,
    )

    aplicar_tabla(
        ws,
        2,
        ws.max_row,
        1,
        ws.max_column,
    )

    for fila in range(
        2,
        ws.max_row + 1,
    ):

        ws.cell(
            fila,
            7,
        ).number_format = "0.00"

        ws.cell(
            fila,
            9,
        ).number_format = "0.00%"

        semaforo_porcentaje(
            ws.cell(
                fila,
                9,
            ),
            ws.cell(
                fila,
                9,
            ).value or 0,
        )

        estado = ws.cell(
            fila,
            10,
        ).value

        if estado == "ALTO":

            ws.cell(
                fila,
                10,
            ).fill = PatternFill(
                "solid",
                fgColor=COLOR_ROJO,
            )

        elif estado == "ATENCI\u00d3N":

            ws.cell(
                fila,
                10,
            ).fill = PatternFill(
                "solid",
                fgColor=COLOR_AMARILLO,
            )

        else:

            ws.cell(
                fila,
                10,
            ).fill = PatternFill(
                "solid",
                fgColor=COLOR_VERDE,
            )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ajustar_columnas(
        ws
    )

    # ---------------------------------------------------------------
    # RANKING
    # ---------------------------------------------------------------

    encabezados_ranking = [
        "Ranking",
        "Unidad",
        "Empresa",
        "Terminal",
        "PPU",
        "Interno",
        "Tipo Bus",
        "Tipo Flota",
        "Plazas",
        "D\u00edas evaluables BD",
        "D\u00edas Sin TX",
        "% Sin TX",
        "Racha m\u00e1xima Sin TX",
        "Primera fecha Sin TX",
        "\u00daltima fecha Sin TX",
        "Fechas Sin TX",
    ]


    def cargar_ranking(
        ws,
        filas,
    ):

        ws.append(
            encabezados_ranking
        )

        for numero, item in enumerate(
            filas,
            start=1,
        ):

            ws.append([
                numero,
                item["unidad"],
                item["empresa"],
                item["terminal"],
                item["ppu"],
                item["interno"],
                item["tipo_bus"],
                item["tipo_flota"],
                item["plazas"],
                item["dias_evaluables"],
                item["dias_sin_tx"],
                item["porcentaje_sin_tx"],
                item["racha_maxima"],
                item["primera_fecha"],
                item["ultima_fecha"],
                item["fechas_txt"],
            ])

        aplicar_encabezado(
            ws,
            1,
            1,
            ws.max_column,
        )

        aplicar_tabla(
            ws,
            2,
            ws.max_row,
            1,
            ws.max_column,
        )

        for fila in range(
            2,
            ws.max_row + 1,
        ):

            ws.cell(
                fila,
                12,
            ).number_format = "0.00%"

            dias = (
                ws.cell(
                    fila,
                    11,
                ).value
                or 0
            )

            semaforo_dias(
                ws.cell(
                    fila,
                    11,
                ),
                dias,
            )

            semaforo_dias(
                ws.cell(
                    fila,
                    13,
                ),
                ws.cell(
                    fila,
                    13,
                ).value or 0,
            )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        ajustar_columnas(
            ws,
            maximo=32,
        )

        ws.column_dimensions[
            "P"
        ].width = 55


    # ---------------------------------------------------------------
    # 3. RANKING GENERAL
    # ---------------------------------------------------------------

    ws = wb.create_sheet(
        "Ranking PPU Sin TX"
    )

    cargar_ranking(
        ws,
        ranking,
    )

    # ---------------------------------------------------------------
    # 4. U8 ALFA
    # ---------------------------------------------------------------

    filas_u8 = [
        item
        for item in ranking
        if item["unidad"] == "U8"
    ]

    ws = wb.create_sheet(
        "U8 ALFA Sin TX"
    )

    cargar_ranking(
        ws,
        filas_u8,
    )

    # ---------------------------------------------------------------
    # 5. U9 OMEGA
    # ---------------------------------------------------------------

    filas_u9 = [
        item
        for item in ranking
        if item["unidad"] == "U9"
    ]

    ws = wb.create_sheet(
        "U9 OMEGA Sin TX"
    )

    cargar_ranking(
        ws,
        filas_u9,
    )

    # ---------------------------------------------------------------
    # 6. DETALLE DIARIO
    # ---------------------------------------------------------------

    ws = wb.create_sheet(
        "Detalle Diario"
    )

    ws.append([
        "Fecha",
        "Unidad",
        "Empresa",
        "Terminal",
        "PPU",
        "Interno",
        "Tipo Bus",
        "Tipo Flota",
        "Plazas",
        "Estado",
    ])

    eventos_ordenados = sorted(
        eventos,
        key=lambda x: (
            x["fecha"],
            x["unidad"],
            x["terminal"],
            x["ppu"],
        ),
    )

    for evento in eventos_ordenados:

        ws.append([
            evento["fecha"],
            evento["unidad"],
            evento["empresa"],
            evento["terminal"],
            evento["ppu"],
            evento["interno"],
            evento["tipo_bus"],
            evento["tipo_flota"],
            evento["plazas"],
            "SIN TX",
        ])

    aplicar_encabezado(
        ws,
        1,
        1,
        ws.max_column,
    )

    aplicar_tabla(
        ws,
        2,
        ws.max_row,
        1,
        ws.max_column,
    )

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ajustar_columnas(
        ws,
        maximo=30,
    )

    # ================================================================
    # SALIDA
    # ================================================================

    buffer = BytesIO()

    wb.save(
        buffer
    )

    buffer.seek(0)

    primera = (
        todas_fechas[0].isoformat()
        if hasattr(
            todas_fechas[0],
            "isoformat"
        )
        else str(
            todas_fechas[0]
        )
    )

    ultima = (
        todas_fechas[-1].isoformat()
        if hasattr(
            todas_fechas[-1],
            "isoformat"
        )
        else str(
            todas_fechas[-1]
        )
    )

    nombre_archivo = (
        "SWAV_HISTORICO_SIN_TX_"
        f"{primera}_"
        f"{ultima}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{nombre_archivo}"'
        },
    )

